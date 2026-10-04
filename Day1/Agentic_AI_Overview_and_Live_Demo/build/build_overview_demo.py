#!/usr/bin/env python3
"""Build 'Agentic AI in action: from concepts to a live Population Health agent' on the SAS EXTERNAL template.

Same design language as Session 2 (copied from 'SAS Viya Agentic AI Experience - Presentation'):
  * SAS 2023 palette (Blue 0766D1, Midnight 032954, Sky 4398F9, Light C4DEFD, Slate 7E889A)
  * Anova Bold / Anova Light theme fonts (embedded in the template)
  * Slate 56pt title + blue 30pt subtitle, light-grey arc band, white rounded cards with soft
    shadows, icon circles overlapping card tops, navy pill headers with blue chevrons,
    slate section dividers, blue title/closing/takeaway slides.

Session: a 45-minute overview of Agentic AI (what it is, RAG, MCP, composite AI) followed by a live
demonstration of the Population Health Agent, and the bridge to the hands-on build.
"""
import json, os, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ICONS = os.path.join(HERE, 'icons')                       # PNG icons rendered by icons.js
ARC_JSON = os.path.join(HERE, 'arc.json')                  # light-grey arc band traced from the reference deck
TEMPLATE = os.environ.get('SAS_TEMPLATE', os.path.normpath(os.path.join(HERE, '..', '..', '..', 'templates', 'SAS_External_Template.pptx')))
OUT = os.environ.get('DEMO_OUT', os.path.normpath(os.path.join(HERE, '..', 'Agentic_AI_Overview_and_Live_Demo.pptx')))

# ---- palette (SAS-2023 theme) --------------------------------------------------------------
BLUE = '0766D1'; NAVY = '032954'; SKY = '4398F9'; LIGHT = 'C4DEFD'; MID = '98C5FB'
SLATE = '7E889A'; SLATE2 = '838E9E'; GRAY = 'BAC0C9'; PANEL = 'E4E7EA'; ARC = 'F0F1F3'
WHITE = 'FFFFFF'; BLACK = '000000'; INK = '262626'; GREEN = '009242'; RED = 'C00000'
PALE = 'EAF3FE'

# ---- page geometry (20 x 11.25 in) ---------------------------------------------------------
X0 = 1.37; XW = 17.25; X1 = X0 + XW; CT = 2.85; CB = 10.4
MJ = '+mj-lt'   # theme major font (Anova Bold)

def rgb(h): return RGBColor.from_string(h)

# ============================================================================================
# low-level helpers (identical to Session 2 so both decks look the same)
# ============================================================================================
def _fill(shape, hexcol):
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(hexcol)

def _no_line(shape):
    shape.line.fill.background()

def _line(shape, hexcol, width=1.0, dash=None):
    shape.line.color.rgb = rgb(hexcol); shape.line.width = Pt(width)
    if dash: shape.line.dash_style = dash

def _effect_clear(spPr):
    for el in spPr.findall(qn('a:effectLst')):
        spPr.remove(el)

def _no_shadow(shape):
    spPr = shape._element.spPr
    _effect_clear(spPr)
    etree.SubElement(spPr, qn('a:effectLst'))

def _shadow(shape, blur=16, dist=3, alpha=16, angle=90):
    spPr = shape._element.spPr
    _effect_clear(spPr)
    eff = etree.SubElement(spPr, qn('a:effectLst'))
    sh = etree.SubElement(eff, qn('a:outerShdw'))
    sh.set('blurRad', str(int(Pt(blur)))); sh.set('dist', str(int(Pt(dist))))
    sh.set('dir', str(angle * 60000)); sh.set('algn', 'ctr'); sh.set('rotWithShape', '0')
    clr = etree.SubElement(sh, qn('a:srgbClr')); clr.set('val', '000000')
    a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(alpha * 1000))

def _bullet(p, color=BLUE, indent=0.26, char='•'):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(Inches(indent)))); pPr.set('indent', str(-int(Inches(indent))))
    for tag in ('a:buClr', 'a:buFont', 'a:buChar', 'a:buNone'):
        for el in pPr.findall(qn(tag)): pPr.remove(el)
    buClr = etree.SubElement(pPr, qn('a:buClr')); c = etree.SubElement(buClr, qn('a:srgbClr')); c.set('val', color)
    buFont = etree.SubElement(pPr, qn('a:buFont')); buFont.set('typeface', 'Arial')
    buChar = etree.SubElement(pPr, qn('a:buChar')); buChar.set('char', char)

_ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
_ANCHOR = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}

def _runs(p, text, size, color, bold, font, italic=False):
    """'**bold**' markup -> Anova Bold runs."""
    parts = text.split('**')
    for i, seg in enumerate(parts):
        if seg == '': continue
        r = p.add_run(); r.text = seg
        f = r.font; f.size = Pt(size); f.color.rgb = rgb(color)
        is_b = bold or (i % 2 == 1)
        if is_b: f.name = MJ
        elif font: f.name = font
        if italic: f.italic = True

def fill_paras(tf, paras, size=18, color=BLACK, bold=False, font=None, align='l',
               line_spacing=None, space_after=None, space_before=None):
    if isinstance(paras, str): paras = [paras]
    first = True
    for item in paras:
        if isinstance(item, str): item = {'t': item}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = _ALIGN[item.get('align', align)]
        ls = item.get('line_spacing', line_spacing)
        if ls: p.line_spacing = ls
        sa = item.get('space_after', space_after)
        if sa is not None: p.space_after = Pt(sa)
        sb = item.get('space_before', space_before)
        if sb is not None: p.space_before = Pt(sb)
        if item.get('bullet'):
            _bullet(p, item.get('bullet_color', BLUE), indent=item.get('indent', 0.26), char=item.get('char', '•'))
        _runs(p, item.get('t', ''), item.get('size', size), item.get('color', color),
              item.get('bold', bold), item.get('font', font), item.get('italic', False))
    return tf

def text(slide, x, y, w, h, paras, size=18, color=BLACK, bold=False, font=None, align='l', anchor='t',
         margins=(0.05, 0.03, 0.05, 0.03), line_spacing=None, space_after=None, rotation=None, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = _ANCHOR[anchor]
    fill_paras(tf, paras, size, color, bold, font, align, line_spacing, space_after)
    if rotation is not None: tb.rotation = rotation
    return tb

def shape_text(shape, paras, size=18, color=WHITE, bold=False, font=None, align='c', anchor='m',
               margins=(0.12, 0.06, 0.12, 0.06), line_spacing=None, space_after=None):
    tf = shape.text_frame; tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = _ANCHOR[anchor]
    fill_paras(tf, paras, size, color, bold, font, align, line_spacing, space_after)
    return shape

def rect(slide, x, y, w, h, fill=WHITE, radius=0.28, shadow=False, line=None, line_w=1.0, dash=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius: s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill: _fill(s, fill)
    else: s.fill.background()
    if line: _line(s, line, line_w, dash)
    else: _no_line(s)
    if shadow: _shadow(s)
    else: _no_shadow(s)
    return s

def card(slide, x, y, w, h, fill=WHITE, radius=0.28, shadow=True, line=None):
    return rect(slide, x, y, w, h, fill=fill, radius=radius, shadow=shadow, line=line)

def pill(slide, x, y, w, h, label, fill=NAVY, color=WHITE, size=16, bold=True, shadow=False):
    s = rect(slide, x, y, w, h, fill=fill, radius=h / 2, shadow=shadow)
    shape_text(s, label, size=size, color=color, bold=bold, margins=(0.1, 0.02, 0.1, 0.02))
    return s

def oval(slide, cx, cy, d, fill=BLUE, shadow=False, line=None, line_w=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    if fill: _fill(s, fill)
    else: s.fill.background()
    if line: _line(s, line, line_w)
    else: _no_line(s)
    if shadow: _shadow(s)
    else: _no_shadow(s)
    return s

def icon(slide, name, variant, x, y, size):
    path = os.path.join(ICONS, f'{name}_{variant}.png')
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(size), Inches(size))

def icon_circle(slide, cx, cy, d, fill, name, variant='w', shadow=True, scale=0.52):
    oval(slide, cx, cy, d, fill=fill, shadow=shadow)
    isz = d * scale
    icon(slide, name, variant, cx - isz / 2, cy - isz / 2, isz)

def arrow(slide, x, y, w, h, fill=BLUE, rotation=0):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(s, fill); _no_line(s); _no_shadow(s)
    if rotation: s.rotation = rotation
    return s

def down_arrow(slide, x, y, w, h, fill=BLUE):
    s = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(s, fill); _no_line(s); _no_shadow(s)
    return s

def chevron(slide, x, y, w, h, fill=BLUE, rotation=0):
    s = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(s, fill); _no_line(s); _no_shadow(s)
    if rotation: s.rotation = rotation
    return s

def connector(slide, x1, y1, x2, y2, color=SLATE, width=1.25, dash=None, head=False, tail=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    if dash: c.line.dash_style = dash
    ln = c.line._get_or_add_ln()
    if head:
        e = etree.SubElement(ln, qn('a:headEnd')); e.set('type', 'triangle'); e.set('w', 'med'); e.set('len', 'med')
    if tail:
        e = etree.SubElement(ln, qn('a:tailEnd')); e.set('type', 'triangle'); e.set('w', 'med'); e.set('len', 'med')
    spPr = c._element.spPr
    _effect_clear(spPr); etree.SubElement(spPr, qn('a:effectLst'))
    return c

def send_to_back(slide, shape, index=2):
    sp = shape._element; tree = slide.shapes._spTree
    tree.remove(sp); tree.insert(index, sp)

ARCPTS = json.load(open(ARC_JSON))['points']
def arc_bg(slide):
    scale = 12700  # EMU per PDF point; reference page = 1440 x 810 pt = 20 x 11.25 in
    fb = slide.shapes.build_freeform(int(ARCPTS[0][0] * scale), int(ARCPTS[0][1] * scale), scale=1.0)
    fb.add_line_segments([(int(x * scale), int(y * scale)) for x, y in ARCPTS[1:]], close=True)
    s = fb.convert_to_shape()
    _fill(s, ARC); _no_line(s); _no_shadow(s)
    send_to_back(slide, s)
    return s

def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt.strip()

def placeholder(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx: return ph
    return None

def drop_placeholder(ph):
    if ph is not None:
        ph._element.getparent().remove(ph._element)

def set_ph(ph, txt, size=None, color=None, bold=None):
    tf = ph.text_frame
    tf.text = txt
    for p in tf.paragraphs:
        for r in p.runs:
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = rgb(color)
            if bold is not None: r.font.bold = bold

# ============================================================================================
# presentation set-up
# ============================================================================================
prs = Presentation(TEMPLATE)
# drop template slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
# keep only master 0 (SAS - EXTERNAL)
mlst = prs.slide_masters._sldMasterIdLst
for el in list(mlst)[1:]:
    prs.part.drop_rel(el.rId); mlst.remove(el)
M = prs.slide_masters[0]
LAY = {l.name: l for l in M.slide_layouts}

def content_slide(title, subtitle=None, arc=True, title_size=None, subtitle_size=30):
    s = prs.slides.add_slide(LAY['SAS - Title & Subtitle'])
    if arc: arc_bg(s)
    t = s.shapes.title
    t.text = title
    if not title_size:
        title_size = 46 if len(title) > 40 else (50 if len(title) > 32 else 56)   # keep titles on one line
    if title_size:
        for p in t.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(title_size)
    ph = placeholder(s, 10)
    if subtitle:
        set_ph(ph, subtitle, size=subtitle_size)
    else:
        drop_placeholder(ph)
    return s

def section_slide(title, subtitle, icon_name, title_size=60):
    s = prs.slides.add_slide(LAY['SAS - Section'])
    set_ph(s.shapes.title, title, size=title_size)
    set_ph(placeholder(s, 10), subtitle, size=36)
    # outlined icon top-right like the reference dividers
    oval(s, 17.15, 2.15, 2.3, fill=None, line=WHITE, line_w=2.5)
    icon(s, icon_name, 'w', 17.15 - 0.62, 2.15 - 0.62, 1.24)
    return s

def heading(slide, x, y, w, txt, size=24, color=INK, h=0.5, align='l'):
    return text(slide, x, y, w, h, txt, size=size, color=color, bold=True, align=align, anchor='t')

def body(slide, x, y, w, h, paras, size=17, color=INK, space_after=6, anchor='t'):
    return text(slide, x, y, w, h, paras, size=size, color=color, space_after=space_after, anchor=anchor)

def icon_row(slide, x, y, w, name, label, desc, circle=BLUE, d=0.7, label_size=18, desc_size=15,
             variant='w', row_h=None, desc_color=INK):
    icon_circle(slide, x + d / 2, y + d / 2, d, circle, name, variant=variant, shadow=False, scale=0.55)
    text(slide, x + d + 0.22, y - 0.06, w - d - 0.22, 0.42, label, size=label_size, color=INK, bold=True)
    tb = text(slide, x + d + 0.22, y + 0.34, w - d - 0.22, (row_h or 1.2) - 0.34, desc, size=desc_size, color=desc_color)
    return tb

def numbered_row(slide, x, y, w, n, label, desc, d=0.62, label_size=18, desc_size=15, row_h=1.0, fill=BLUE):
    c = oval(slide, x + d / 2, y + d / 2, d, fill=fill)
    shape_text(c, str(n), size=18, color=WHITE, bold=True, margins=(0, 0, 0, 0))
    text(slide, x + d + 0.22, y - 0.05, w - d - 0.22, 0.42, label, size=label_size, color=INK, bold=True)
    if desc:
        text(slide, x + d + 0.22, y + 0.34, w - d - 0.22, row_h - 0.34, desc, size=desc_size, color=INK)

def callout(slide, x, y, w, h, paras, fill=LIGHT, color=NAVY, size=18, icon_name=None, icon_variant='n', align='l'):
    r = rect(slide, x, y, w, h, fill=fill, radius=0.22)
    pad = 0.35
    if icon_name:
        isz = min(0.7, h - 0.4)
        icon(slide, icon_name, icon_variant, x + 0.35, y + (h - isz) / 2, isz)
        pad = 0.35 + isz + 0.3
    text(slide, x + pad, y + 0.12, w - pad - 0.3, h - 0.24, paras, size=size, color=color, anchor='m', align=align)
    return r

def copyright_line(slide, color=LIGHT):
    text(slide, 1.37, 10.78, 6.0, 0.27, 'Copyright © SAS Institute Inc. All rights reserved.', size=10, color=color, anchor='b')

def top_card(slide, x, y, w, h, icon_name, circle, title, desc, foot=None, title_size=20, desc_size=15,
             fill=WHITE, title_color=INK, desc_color=INK, foot_color=BLUE, icon_d=1.15, desc_h=None, icon_variant='w'):
    """White card with an icon circle overlapping its top edge, a centred title, body text and an optional blue foot line."""
    card(slide, x, y, w, h, fill=fill)
    icon_circle(slide, x + w / 2, y, icon_d, circle, icon_name, variant=icon_variant)
    text(slide, x + 0.3, y + icon_d * 0.72, w - 0.6, 0.65, title, size=title_size, color=title_color, bold=True, align='c')
    text(slide, x + 0.35, y + icon_d * 0.72 + 0.75, w - 0.7, desc_h or (h - icon_d * 0.72 - 0.75 - (1.2 if foot else 0.3)),
         desc, size=desc_size, color=desc_color)
    if foot:
        text(slide, x + 0.35, y + h - 1.15, w - 0.7, 0.95, [{'t': foot, 'color': foot_color, 'size': desc_size - 1, 'bold': True}], anchor='b')

# ============================================================================================
# SLIDES
# ============================================================================================
# 1 ---- Title -------------------------------------------------------------------------------
s = prs.slides.add_slide(LAY['SAS - Title'])
set_ph(s.shapes.title, 'Agentic AI in action', size=66)
set_ph(placeholder(s, 10), 'From the concepts to a live Population Health agent', size=36)
set_ph(placeholder(s, 11), 'Raed Aldweik  |  SAS\nEHS Agentic AI Bootcamp', size=24)
notes(s, """
Welcome. This session does two things in 45 minutes: it gives the room a shared vocabulary for Agentic AI (what an agent
is, what RAG, MCP and composite AI add to it), and it shows a real agent at work on a population-health problem, the
same agent the participants will build themselves in the hands-on part of the bootcamp. No code, no maths; every
concept is shown once in plain words and once live in the demo.
Timing guide: intro 3 min (slides 1-2), what is Agentic AI 8 min (3-6), the building blocks 10 min (7-12), live demo
18 min (13-18, of which about 12 min live in SAS Retrieval Agent Manager), what you will build 4 min (19-21), wrap-up
2 min (22-23). About 45 minutes. If you are short on time, skip slide 6 (healthcare examples) and slide 21 (the
component table). If the live environment misbehaves, slide 17 is the backup walk-through of one answer.
Before the session: sign in to RAM, open the Population Health Agent, run question 1 once so the compute session is
warm, and keep the app's SAS RAM tab open on the second screen for the trace.
""")

# 2 ---- Agenda ------------------------------------------------------------------------------
s = content_slide('In the next 45 minutes', 'Four parts: the concepts, the building blocks, a live agent, and what you will build')
items = [
    ('What Agentic AI is, and is not', 'From predicting and generating to acting: the agent loop, and where agents fit in healthcare'),
    ('The building blocks: LLM, RAG, MCP, composite AI', 'What each one does for the agent, in plain words, and why the agent needs all of them'),
    ('Live demo: the Population Health Agent', 'One registry, four national guidelines, one risk model, and the trace behind every answer'),
    ('What you will build in this bootcamp', 'The same agent, assembled by you in SAS Viya and SAS Retrieval Agent Manager, without code'),
]
connector(s, 1.82, 3.4, 1.82, 8.75, color=LIGHT, width=2.5)
for i, (a, b) in enumerate(items):
    y = 3.05 + i * 1.72
    sq = rect(s, 1.37, y, 0.9, 0.9, fill=BLUE, radius=0.12)
    shape_text(sq, str(i + 1), size=26, color=WHITE, bold=True, margins=(0, 0, 0, 0))
    text(s, 2.55, y - 0.04, 8.3, 0.5, a, size=22, color=INK, bold=True)
    text(s, 2.55, y + 0.44, 8.3, 0.85, b, size=15, color=SLATE)
card(s, 11.4, 2.95, 7.22, 7.0)
heading(s, 11.85, 3.3, 6.4, 'By the end of this session you can…', size=22, color=BLUE)
outs = [
    ('Explain Agentic AI in one sentence', 'and tell an agent apart from a chatbot and from a script'),
    ('Say what RAG, MCP and composite AI do', 'for an agent, in plain words, and why it needs all three'),
    ('Read the Population Health Agent', 'its data, documents, model, tools and guardrails, before you build your own'),
]
for i, (a, b) in enumerate(outs):
    y = 4.25 + i * 1.75
    icon_row(s, 11.85, y, 6.4, 'check', a, b, circle=NAVY, d=0.62, label_size=17, desc_size=15, row_h=1.4)
notes(s, """
Walk the four parts quickly (about 8, 10, 18 and 4 minutes). Say up front that the demo agent and the agent they will
build are the same agent: same registry, same guidelines, same model, same tools. Everything in parts 1 and 2 is there
so that the demo makes sense, and everything in the demo is there so that the build makes sense. The card on the right
is the contract with the room: three things they will be able to do.
""")

# 3 ---- Section: what is agentic AI --------------------------------------------------------
s = section_slide('What is Agentic AI?', 'From answering questions to getting work done', 'robot')
notes(s, "Part 1 of 4. About 8 minutes: three generations of AI, the agent loop, and where agents fit in healthcare.")

# 4 ---- Three generations -------------------------------------------------------------------
s = content_slide('Predict, generate, act', 'Agentic AI is the third step, and it stands on the other two')
arrow(s, X0, 2.95, XW, 0.6, fill=LIGHT)
text(s, X0 + 0.4, 2.95, XW - 1.2, 0.6, 'Predict  →  Generate  →  Act', size=16, color=NAVY, bold=True, anchor='m')
gens = [
    ('model', 'Analytical AI and machine learning', 'Learns from data to predict and classify: a deterioration risk score, a demand forecast, a fraud flag, a no-show probability.',
     'In the demo: the deterioration-risk model behind the agent.', 'Answers: how likely? how many? which ones?', NAVY, WHITE, INK, BLUE),
    ('sparkles', 'Generative AI', 'Learns language to create: it summarises a guideline, drafts a letter, explains a chart, answers one question at a time.',
     'In the demo: the language model that writes every answer.', 'Answers: what does this say? write this for me.', SKY, WHITE, INK, BLUE),
    ('robot', 'Agentic AI', 'Pursues a goal: breaks it into steps, uses tools, models and documents, checks the results, and keeps going until the job is done, inside boundaries set by people.',
     'In the demo: the Population Health Agent itself.', 'Answers: get this done, and show me how you did it.', BLUE, BLUE, WHITE, LIGHT),
]
gw3 = (XW - 2 * 0.45) / 3
for i, (ic, t, d, ex, foot, circ, fill, tc, fc) in enumerate(gens):
    x = X0 + i * (gw3 + 0.45); y = 4.1; h = 4.85
    top_card(s, x, y, gw3, h, ic, circ if fill == WHITE else WHITE, t, [d, {'t': ex, 'space_before': 12, 'color': fc, 'bold': True}],
             foot=foot, fill=fill, title_color=tc, desc_color=tc, foot_color=fc, title_size=21, desc_size=16, icon_variant=('w' if fill == WHITE else 'b'))
callout(s, X0, 9.2, XW, 1.0, 'Agentic AI does not replace the first two. It uses them as tools: the model scores, the LLM reasons and writes, the agent decides what to do next.', icon_name='lightbulb', size=16)
notes(s, """
Three generations, one sentence each. Analytical AI predicts: EHS already runs this (risk scores, forecasts). Generative
AI creates language: most of the room has used a chatbot. Agentic AI acts: it is given a goal, and it works through the
steps itself, calling the other two as tools. The callout is the point of the whole session: the agent is not a better
chatbot, it is an orchestrator of things you already have. The third card is blue on purpose: it is the one we build.
""")

# 5 ---- The loop ----------------------------------------------------------------------------
s = content_slide('What makes it an agent: the loop', 'A chatbot answers once. An agent keeps going until the goal is met')
lcx, lcy, R = 5.95, 6.75, 2.35
ring = s.shapes.add_shape(MSO_SHAPE.DONUT, Inches(lcx - R - 0.42), Inches(lcy - R - 0.42), Inches(2 * R + 0.84), Inches(2 * R + 0.84))
ring.adjustments[0] = 0.15; _fill(ring, LIGHT); _no_line(ring); _no_shadow(ring)
nodes = [('PERCEIVE', 270, NAVY), ('REASON & PLAN', 0, BLUE), ('ACT', 90, NAVY), ('OBSERVE', 180, BLUE)]
for lab, ang, col in nodes:
    a = math.radians(ang); nx = lcx + R * math.cos(a); ny = lcy + R * math.sin(a)
    o = oval(s, nx, ny, 1.6, fill=col, shadow=True)
    shape_text(o, lab, size=13, color=WHITE, bold=True, margins=(0.05, 0.02, 0.05, 0.02))
for phi in (315, 45, 135, 225):
    a = math.radians(phi); px_ = lcx + R * math.cos(a); py_ = lcy + R * math.sin(a)
    chevron(s, px_ - 0.3, py_ - 0.26, 0.6, 0.52, fill=BLUE, rotation=phi + 90)
text(s, lcx - 1.35, lcy - 0.85, 2.7, 1.7, ['Repeat until the goal is met,', {'t': 'or a boundary is hit', 'color': BLUE}], size=14, color=NAVY, bold=True, align='c', anchor='m')
rx = 10.0; rw = X1 - rx
card(s, rx, 3.0, rw, 5.35)
heading(s, rx + 0.45, 3.3, rw - 0.9, 'Each turn of the loop', size=20, color=BLUE)
lrows = [('eye', 'Perceive', 'Read the situation: the question, the data that came back, the last tool result.'),
         ('brain', 'Reason & plan', 'The language model weighs goal, rules, knowledge and tools, and proposes the next step.'),
         ('play', 'Act', 'The runtime executes it: a query, a model score, a document search, a draft.'),
         ('search', 'Observe', 'The result comes back and becomes the next thing to perceive.')]
for i, (ic, a, b) in enumerate(lrows):
    icon_row(s, rx + 0.45, 3.95 + i * 1.08, rw - 0.9, ic, a, b, circle=NAVY if i % 2 == 0 else BLUE, d=0.58, label_size=16, desc_size=14, row_h=1.0)
callout(s, rx, 8.65, rw, 1.55, 'An AI agent is a system that pursues a **goal** by reasoning about what to do next, using **tools** to act, '
        'and checking the result, inside **boundaries** set by people.', icon_name='robot', size=15)
notes(s, """
This loop is the definition of 'agentic'. Four beats: perceive, reason and plan, act, observe, then round again. A
chatbot runs the loop once; an agent runs it until the goal is met or a boundary stops it. Three words carry the rest
of the session: goal, tools, boundaries. If people ask 'where is the LLM?', point to Reason & plan: that is where it
lives, and only there. Everything else in the loop is ordinary software that we control. You will see this loop in the
demo as the trace: every turn shows up as a retrieval, a query or a model call.
""")

# 6 ---- Where agents fit in healthcare -----------------------------------------------------
s = content_slide('Where agents fit in healthcare', 'Real-world patterns: the agent reads, reasons and drafts; a person decides')
uses = [
    ('heart', 'Population health', 'Finds cohorts and care gaps in the registry, scores risk, drafts a recall list with the guideline section that justifies it.', 'The programme lead approves the list.', True),
    ('calendar', 'Outpatient no-shows', 'Spots who is likely to miss tomorrow\'s appointment, confirms or reschedules, and refills the freed slot from the waiting list.', 'The coordinator handles exceptions.', False),
    ('clipboard', 'Coding and documentation', 'Reads the discharge note, proposes the codes and a summary, with the supporting sentences highlighted.', 'The coder signs off.', False),
    ('chat', 'Patient enquiries', 'Reads messages and complaints, routes them to the right department, drafts the reply in the patient\'s language.', 'A person sends anything clinical.', False),
    ('trend', 'Capacity and supply', 'Forecasts demand per facility, flags bed pressure and stock-outs, proposes transfers and orders.', 'Operations decides the move.', False),
    ('badge', 'Quality and audit reporting', 'Computes the quality measures, explains the drivers, drafts the report against the national standard.', 'The quality team reviews and publishes.', False),
]
gw = (XW - 2 * 0.4) / 3; gh = 3.3
for i, (ic, t, d, foot, demo) in enumerate(uses):
    r, c = divmod(i, 3)
    x = X0 + c * (gw + 0.4); y = 2.95 + r * (gh + 0.4)
    card(s, x, y, gw, gh)
    icon_circle(s, x + 0.65, y + 0.62, 0.72, BLUE if demo else NAVY, ic, shadow=False, scale=0.55)
    text(s, x + 1.2, y + 0.35, gw - 1.45 - (1.7 if demo else 0), 0.55, t, size=19, color=INK, bold=True)
    if demo:
        pill(s, x + gw - 1.95, y + 0.4, 1.65, 0.46, 'Today\'s demo', fill=BLUE, size=12)
    text(s, x + 0.4, y + 1.15, gw - 0.8, 1.45, d, size=14, color=INK)
    text(s, x + 0.4, y + gh - 0.75, gw - 0.8, 0.5, [{'t': foot, 'color': BLUE, 'size': 13, 'bold': True}], anchor='b')
notes(s, """
Six patterns, all real, all the same shape: the agent reads, reasons and drafts; a person decides. Read the blue foot
line of each card aloud, because that is what makes them acceptable in a hospital. The first card is today's demo and
the bootcamp's build. Invite the room to map their own process onto one of the six; most back-office and
programme-management work fits. Skip this slide if you are behind.
""")

# 7 ---- Section: building blocks -----------------------------------------------------------
s = section_slide('The building blocks', 'LLM, RAG, MCP and composite AI, in plain words', 'blocks')
notes(s, "Part 2 of 4. About 10 minutes: the map, then one slide each for the language model, RAG, MCP, and composite AI with one question flowing through all of them.")

# 8 ---- Map of the blocks -------------------------------------------------------------------
s = content_slide('Four blocks around one reasoning engine', 'The LLM reasons; everything that makes it useful and safe sits around it')
cx, cy = 10.0, 6.55
blocks = [
    ('files', 'Knowledge  ·  RAG', 'Your documents, retrieved on demand and cited: the national guidelines, policies, protocols.'),
    ('api', 'Tools  ·  MCP', 'A standard plug into data, models and systems: the agent asks, the server executes, every call is logged.'),
    ('model', 'Analytics  ·  composite AI', 'The ML models, SQL queries and decision rules you already trust, called as tools.'),
    ('oversight', 'People  ·  human oversight', 'Goal, boundaries, approvals and review. The agent drafts; a person decides.'),
]
cw_, ch_ = 6.2, 2.55
pos = [(X0, 3.1), (X0, 7.0), (X1 - cw_, 3.1), (X1 - cw_, 7.0)]
for (ic, lab, desc), (x, y) in zip(blocks, pos):
    left = x < 9
    connector(s, (x + cw_) if left else x, y + ch_ / 2, cx + (-1.55 if left else 1.55), cy, color=MID, width=1.5, dash=MSO_LINE.ROUND_DOT)
    card(s, x, y, cw_, ch_)
    icon_circle(s, x + 0.95, y + ch_ / 2, 1.1, NAVY if ic == 'oversight' else BLUE, ic, shadow=False)
    text(s, x + 1.75, y + 0.35, cw_ - 2.0, 0.5, lab, size=21, color=INK, bold=True)
    text(s, x + 1.75, y + 0.95, cw_ - 2.0, 1.4, desc, size=15, color=INK)
oval(s, cx, cy, 3.1, fill=NAVY, shadow=True)
oval(s, cx, cy, 2.75, fill=None, line=WHITE, line_w=1.5)
icon(s, 'brain', 'w', cx - 0.4, cy - 1.05, 0.8)
text(s, cx - 1.3, cy - 0.2, 2.6, 1.2, ['LLM', {'t': 'reasons, plans, writes', 'size': 13}], size=22, color=WHITE, bold=True, align='c', anchor='t')
text(s, X0, 9.85, XW, 0.4, 'One slide for each block follows. The language model never touches data or systems directly; the blocks around it do.', size=13, color=SLATE)
notes(s, """
The map for the next four slides. The language model is in the middle because it reasons, plans and writes, and only
that. Knowledge is what it can quote (RAG). Tools are what it can call (MCP). Models and rules are the analytics you
already own, exposed as tools (composite AI). People set the goal and the boundaries and take the decisions. Say the
footer sentence slowly: the model never touches your data or systems directly. That is the sentence that makes the
security people in the room relax.
""")

# 9 ---- LLM ---------------------------------------------------------------------------------
s = content_slide('The LLM: a reasoning engine with no access', 'What the language model is good at, and what it cannot know on its own')
hw = (XW - 0.45) / 2
card(s, X0, 3.0, hw, 5.75)
icon_circle(s, X0 + 0.7, 3.65, 0.8, GREEN, 'check', shadow=False, scale=0.5)
text(s, X0 + 1.25, 3.4, hw - 1.5, 0.5, 'Good at', size=20, color=INK, bold=True)
text(s, X0 + 0.45, 4.45, hw - 0.9, 3.2, [
    {'t': 'Understanding the question, in English or Arabic', 'bullet': True},
    {'t': 'Breaking a goal into steps and choosing the right tool for each', 'bullet': True},
    {'t': 'Reading a retrieved passage and summarising it faithfully', 'bullet': True},
    {'t': 'Writing the answer: clear, structured, the cohort stated in words', 'bullet': True},
], size=18, color=INK, space_after=12)
text(s, X0 + 0.45, 7.7, hw - 0.9, 0.85, [{'t': 'In the demo: it picks the SQL tool for a count and the collection for a target.', 'color': BLUE, 'size': 14, 'bold': True}], anchor='b')
gx2 = X0 + hw + 0.45
card(s, gx2, 3.0, hw, 5.75)
icon_circle(s, gx2 + 0.7, 3.65, 0.8, RED, 'x', shadow=False, scale=0.5)
text(s, gx2 + 1.25, 3.4, hw - 1.5, 0.5, 'Cannot, on its own', size=20, color=INK, bold=True)
text(s, gx2 + 0.45, 4.45, hw - 0.9, 3.2, [
    {'t': 'Know your registry: how many patients are uncontrolled today', 'bullet': True, 'bullet_color': RED},
    {'t': 'Know the current national guideline, section by section', 'bullet': True, 'bullet_color': RED},
    {'t': 'Do reliable arithmetic across 4,000 rows', 'bullet': True, 'bullet_color': RED},
    {'t': 'Run anything: it proposes an action, it never executes one', 'bullet': True, 'bullet_color': RED},
], size=18, color=INK, space_after=12)
text(s, gx2 + 0.45, 7.7, hw - 0.9, 0.85, [{'t': 'In the demo: every number is a query that ran, every target is a citation.', 'color': BLUE, 'size': 14, 'bold': True}], anchor='b')
callout(s, X0, 9.15, XW, 1.05, 'That is why we surround it: **RAG** for the words it must quote, **tools and models** for the numbers it must not guess, '
        '**people** for the decisions it must not take.', icon_name='lightbulb', size=16)
notes(s, """
Give the language model its due and its limits. Good at language, planning and faithful summarising. Cannot know your
data or your current guideline, cannot count reliably, cannot run anything. When none of its sources contains the
answer it produces a plausible guess: that is the hallucination problem, and the cure is design, not hope. The callout
is the structure of the next three slides: RAG, tools and models, people. Ask the room which of the four 'cannot'
items worries them most; the answer is usually the guideline, which is exactly what RAG fixes.
""")

# 10 ---- RAG --------------------------------------------------------------------------------
s = content_slide('RAG: Retrieval-Augmented Generation', 'Answers from your documents, with the section cited, or “not covered”')
def rag_row(y, h, label, steps):
    rect(s, X0, y, XW, h, fill=PANEL, radius=0.3, shadow=True)
    p = rect(s, X0 + 0.45, y - 0.3, 5.6, 0.6, fill=NAVY, radius=0.3, shadow=True)
    shape_text(p, label, size=15, color=WHITE, bold=True, margins=(0.1, 0.02, 0.1, 0.02))
    n = len(steps); gap = 0.62; cw = (XW - 0.9 - (n - 1) * gap) / n
    for i, (ic, t, d) in enumerate(steps):
        x = X0 + 0.45 + i * (cw + gap); yy = y + 0.5; hh = h - 0.8
        card(s, x, yy, cw, hh, shadow=False)
        icon_circle(s, x + 0.62, yy + 0.58, 0.68, BLUE if i % 2 == 0 else NAVY, ic, shadow=False, scale=0.55)
        text(s, x + 1.1, yy + 0.3, cw - 1.3, 0.6, t, size=16, color=INK, bold=True)
        text(s, x + 0.3, yy + 1.05, cw - 0.6, hh - 1.15, d, size=13, color=INK)
        if i < n - 1:
            arrow(s, x + cw + 0.1, yy + hh / 2 - 0.2, gap - 0.2, 0.4, fill=BLUE)
rag_row(3.2, 3.3, 'Before: build the knowledge base', [
    ('files', 'Documents', 'The four NHA guideline PDFs: diabetes, lipids, hypertension, screening and recall.'),
    ('split', 'Chunk', 'Split into passages of 600 to 800 characters, with overlap, so a section stays together.'),
    ('embed', 'Embedding model', 'Turn each passage into a vector: a numeric fingerprint of its meaning.'),
    ('vector', 'Vector store', 'Keep the vectors so that passages can be found by meaning, not by keyword.'),
])
rag_row(6.85, 3.3, 'At question time: retrieve, then generate', [
    ('question', 'Question', '“What is the LDL target for a very-high-risk patient?”'),
    ('search', 'Retrieve', 'Embed the question; fetch the 4 to 6 passages closest in meaning.'),
    ('brain', 'Generate', 'The LLM writes the answer from those passages only, nothing from memory.'),
    ('cite', 'Cited answer', '“Below 1.4 mmol/L, NHA-CG-02 §3.” Or: “not covered by the guidelines in my collection.”'),
])
notes(s, """
RAG in two rows. Row one happens once, when you build the knowledge base: documents are chunked, each chunk is turned
into a vector (a fingerprint of its meaning) and stored. Row two happens at every question: the question is turned
into a vector, the closest passages are fetched, and the model writes from those passages only, with the citation.
The last card matters most for healthcare: when the collection does not cover the question, the agent says so instead
of guessing. In SAS this whole slide is a Retrieval Agent Manager 'collection'; the participants create one from the
four PDFs in the hands-on part, and in the demo you will see the retrieved passages in the trace.
""")

# 11 ---- MCP --------------------------------------------------------------------------------
s = content_slide('MCP: the Model Context Protocol', 'One standard plug between the agent and your data, models and systems')
ax, ay, aw_, ah = X0, 3.3, 4.6, 5.3
card(s, ax, ay, aw_, ah)
icon_circle(s, ax + aw_ / 2, ay, 1.15, NAVY, 'robot')
text(s, ax + 0.3, ay + 0.8, aw_ - 0.6, 0.6, 'The agent', size=22, color=INK, bold=True, align='c')
arows = [('goal', 'Goal & rules'), ('brain', 'Reasoning engine (LLM)'), ('files', 'Knowledge (RAG)'), ('memory', 'Memory of this task')]
for i, (ic, t) in enumerate(arows):
    yy = ay + 1.65 + i * 0.88
    icon_circle(s, ax + 0.75, yy + 0.3, 0.6, BLUE, ic, shadow=False, scale=0.55)
    text(s, ax + 1.25, yy + 0.05, aw_ - 1.5, 0.5, t, size=16, color=INK, bold=True)
# MCP server block
gx, gy, gw2, gh2 = 7.3, 3.0, 2.5, 5.9
rect(s, gx, gy, gw2, gh2, fill=NAVY, radius=0.3, shadow=True)
text(s, gx - 2.3, gy + gh2 / 2 - 0.9, gh2 - 0.4, 1.8, ['MCP SERVER', {'t': 'tool descriptions  ·  permissions  ·  audit log', 'size': 13, 'bold': False}],
     size=18, color=WHITE, bold=True, align='c', anchor='m', rotation=270)
tb = s.shapes[-1]; tb.left = Inches(gx + gw2 / 2 - (gh2 - 0.4) / 2); tb.top = Inches(gy + gh2 / 2 - 0.9)
connector(s, ax + aw_ + 0.1, 5.95, gx - 0.1, 5.95, color=BLUE, width=2.25, tail=True, head=True)
text(s, ax + aw_ + 0.1, 5.25, gx - ax - aw_ - 0.2, 0.6, 'asks  ·  gets results', size=12, color=BLUE, bold=True, align='c', anchor='b')
# tools stack
tools_ = [('sql', 'Query data (SQL)', 'FedSQL over the registry tables in SAS Viya CAS: counts, rates, costs, one patient'),
          ('model', 'Score the ML model', 'The published deterioration-risk model in SAS Micro Analytic Service'),
          ('rules', 'Run a decision flow', 'SAS Intelligent Decisioning rules: recall priority, the 30-day window'),
          ('chart', 'Generate a chart', 'Control by emirate, cost by risk tier, straight into the answer')]
sx = 10.9; sw = X1 - sx; sh = 1.24; sgap = 0.31
for i, (ic, t, d) in enumerate(tools_):
    yy = gy + i * (sh + sgap)
    card(s, sx, yy, sw, sh)
    icon_circle(s, sx + 0.7, yy + sh / 2, 0.72, BLUE if i % 2 == 0 else SKY, ic, shadow=False, scale=0.55)
    text(s, sx + 1.25, yy + 0.12, sw - 1.45, 0.45, t, size=16, color=INK, bold=True)
    text(s, sx + 1.25, yy + 0.55, sw - 1.45, 0.65, d, size=12, color=SLATE)
    connector(s, gx + gw2 + 0.05, yy + sh / 2, sx - 0.08, yy + sh / 2, color=BLUE, width=1.75, tail=True)
text(s, sx, gy - 0.5, sw, 0.4, 'Tools, described in plain language', size=14, color=BLUE, bold=True)
callout(s, X0, 9.2, XW, 1.0, 'Think USB-C: one plug, any device. Tools are described in plain language so you can read what the agent can reach; '
        'the agent asks and the server executes with the permissions you granted; every call is logged.', icon_name='usb', size=15)
notes(s, """
MCP is the standard that lets any agent talk to any tool through one plug: the USB-C of agents. The agent on the left
never reaches a database or a model directly. It sends a request ('score patient EHS-100092') to the MCP server, which
knows what the agent is allowed to call, executes it, logs it, and returns the result. Three things for a non-technical
audience: tools are described in plain language, so a business owner can read the tool list; the model asks and the
runtime executes, which is where permissions live; and everything is logged, which is the trace you will see in the
demo. In the bootcamp each team gets an MCP server scoped to its own table and model.
""")

# 12 ---- Composite AI -----------------------------------------------------------------------
s = content_slide('Composite AI: one question, several engines', 'Each kind of fact comes from the engine that is good at it')
qx, qw = X0, 4.4
card(s, qx, 3.4, qw, 4.6)
icon_circle(s, qx + qw / 2, 3.4, 1.1, NAVY, 'question')
text(s, qx + 0.3, 4.2, qw - 0.6, 0.5, 'The question', size=18, color=INK, bold=True, align='c')
text(s, qx + 0.4, 4.85, qw - 0.8, 2.9, '“Who would qualify for a GLP-1 RA under the guideline, and how many are already on one?”', size=17, color=NAVY, align='c', anchor='m')
ex, ew = 6.45, 6.7
engines = [('files', 'Guideline retrieval (RAG)', 'Finds NHA-CG-01 §4: consider a GLP-1 RA when HbA1c stays above target and BMI > 30.', BLUE),
           ('sql', 'SQL on the registry', 'Counts the eligible cohort, then the share not yet treated. Every number is a query that ran.', NAVY),
           ('model', 'Machine learning', 'Scores the cohort with the deterioration model so the clinic reviews the riskiest first.', SKY),
           ('rules', 'Decision rules', 'Recall within 30 days; consent exclusions; the clinician signs off (NHA-CG-01 §7).', BLUE)]
eh = 1.32; egap = 0.25
for i, (ic, t, d, col) in enumerate(engines):
    yy = 2.95 + i * (eh + egap)
    card(s, ex, yy, ew, eh)
    icon_circle(s, ex + 0.7, yy + eh / 2, 0.74, col, ic, shadow=False, scale=0.55)
    text(s, ex + 1.3, yy + 0.1, ew - 1.5, 0.45, t, size=16, color=INK, bold=True)
    text(s, ex + 1.3, yy + 0.55, ew - 1.5, 0.75, d, size=12, color=INK)
arrow(s, qx + qw + 0.15, 5.5, ex - qx - qw - 0.3, 0.5, fill=BLUE)
anx = ex + ew + 0.95; anw = X1 - anx
arrow(s, ex + ew + 0.15, 5.5, anx - ex - ew - 0.3, 0.5, fill=BLUE)
card(s, anx, 2.95, anw, 6.25, fill=PALE, shadow=True)
icon_circle(s, anx + anw / 2, 2.95, 1.1, BLUE, 'cite')
text(s, anx + 0.3, 3.75, anw - 0.6, 0.5, 'The answer', size=18, color=INK, bold=True, align='c')
text(s, anx + 0.35, 4.35, anw - 0.7, 4.7, [
    {'t': 'Type 2, HbA1c ≥ 9%, BMI > 30: **399 patients**', 'bullet': True},
    {'t': 'Of whom **373** are not yet on an SGLT2i or GLP-1 RA', 'bullet': True},
    {'t': 'Guideline: consider a GLP-1 RA, **NHA-CG-01 §4**', 'bullet': True},
    {'t': 'Top ten by model risk for the clinic to review first', 'bullet': True},
    {'t': 'A draft list; **the clinician decides** (§7)', 'bullet': True},
    {'t': 'Source line: the query, the sections used', 'bullet': True},
], size=14, color=NAVY, space_after=7)
callout(s, X0, 9.45, XW, 0.8, 'The LLM never computes a number or invents a threshold. It decides which engine to call, and writes the answer from what came back.', icon_name='puzzle', size=15)
notes(s, """
Composite AI is the reason the agent is trustworthy: each kind of fact comes from the engine that is good at it. One
real question from the demo flows left to right. The retrieval engine finds the section; SQL counts the cohort; the
model ranks it; the rules say who may see what and who decides; the language model only orchestrates and writes.
Read the answer card: every number has a cohort, every clinical statement has a section, and the last line hands the
decision to the clinician. This is the slide to come back to when someone asks 'but can we trust what it says?'.
""")

# 13 ---- Section: live demo ------------------------------------------------------------------
s = section_slide('Live demo: the Population Health Agent', 'One registry, four national guidelines, one risk model, one agent', 'stethoscope', title_size=54)
notes(s, "Part 3 of 4. About 18 minutes: the scenario and the solution overview (3 min), the demo plan (1 min), the live demo in SAS Retrieval Agent Manager (12 min), then the backup walk-through and the guardrails (2 min).")

# 14 ---- Scenario ---------------------------------------------------------------------------
s = content_slide('The scenario: diabetes across five emirates', 'One programme lead, 4,000 patients, 18 facilities, and the questions asked every week')
sc = [('database', 'Structured data', 'A diabetes registry modelled on the national health information exchange: one row per patient, 4,000 patients across 18 hospitals and health centres in 5 emirates, 54 columns: HbA1c, blood pressure, lipids, kidneys, therapy, open care gaps, utilisation and cost.', 'Loaded as a CAS table in SAS Viya', NAVY),
      ('files', 'Unstructured data', 'Four national guidelines, section-numbered, as PDFs: NHA-CG-01 type 2 diabetes management, NHA-CG-02 cardiovascular risk and lipids, NHA-CG-03 hypertension, NHA-PP-01 population screening and recall.', 'Indexed as a collection in SAS Retrieval Agent Manager', BLUE),
      ('model', 'The machine learning model', 'Predicts a deterioration event in the next 12 months: an admission for hypo- or hyperglycaemia, DKA, a foot infection, acute kidney injury, or HbA1c rising to 10% or more. Event rate in the registry: 10.4%.', 'Champion from Model Studio, published for scoring', SKY)]
scw = (XW - 2 * 0.45) / 3
for i, (ic, t, d, foot, col) in enumerate(sc):
    x = X0 + i * (scw + 0.45); y = 3.5; h = 4.75
    top_card(s, x, y, scw, h, ic, col, t, d, foot=foot, title_size=21, desc_size=14)
stats = [('4,000', 'patients in the registry'), ('34.9%', 'well controlled (HbA1c below 7%)'), ('AED 64.0M', 'annual cost, last 12 months'), ('10.4%', 'had a deterioration event in 12 months')]
stw = (XW - 3 * 0.4) / 4
for i, (n, lab) in enumerate(stats):
    x = X0 + i * (stw + 0.4); y = 8.55
    rect(s, x, y, stw, 1.0, fill=PALE, radius=0.2)
    text(s, x + 0.25, y + 0.08, stw - 0.5, 0.5, n, size=24, color=NAVY, bold=True)
    text(s, x + 0.25, y + 0.58, stw - 0.5, 0.38, lab, size=12, color=SLATE)
text(s, X0, 9.75, XW, 0.45, 'All patient data is synthetic and the guidelines are a teaching set written for this bootcamp. Not for clinical use.', size=13, color=SLATE)
notes(s, """
Set the scene in 90 seconds. The programme lead runs the diabetes registry across the Northern Emirates and asks the
same questions every week: how many are uncontrolled, where are the gaps, who is at risk, what does the guideline
say, what would a programme intervention save. Three ingredients, one per card: the registry (structured, in CAS),
the guidelines (unstructured, in a RAM collection), and the model (built in Model Studio, published for scoring).
Say clearly that the data is synthetic and the guidelines are a teaching set: the numbers are real for this data, not
for EHS. These three cards are exactly the three things the participants will create.
""")

# 15 ---- Solution overview ------------------------------------------------------------------
s = content_slide('Solution overview', 'How the pieces fit: data, agent, MCP server and tools')
# -- left column: unstructured data
text(s, X0, 2.9, 4.3, 0.4, 'Unstructured data', size=16, color=BLUE, bold=True, align='c')
rect(s, X0, 3.3, 4.3, 3.2, fill=None, radius=0.25, line=BLUE, line_w=1.25, dash=MSO_LINE.ROUND_DOT)
d1 = rect(s, X0 + 0.3, 3.5, 3.7, 0.7, fill=PANEL, radius=0.12)
shape_text(d1, 'Clinical guidelines & policy documents', size=13, color=INK, bold=True)
text(s, X0 + 0.3, 4.22, 3.7, 0.5, 'National clinical guidelines, care protocols, policy frameworks', size=11, color=SLATE, align='c')
down_arrow(s, X0 + 2.15 - 0.17, 4.72, 0.34, 0.32, fill=BLUE)
e1 = rect(s, X0 + 0.6, 5.08, 3.1, 0.52, fill=LIGHT, radius=0.1)
shape_text(e1, 'Embedding model', size=13, color=NAVY, bold=True)
down_arrow(s, X0 + 2.15 - 0.17, 5.64, 0.34, 0.3, fill=BLUE)
v1 = rect(s, X0 + 0.3, 5.96, 3.7, 0.5, fill=BLUE, radius=0.1)
shape_text(v1, 'Vector store (RAG)', size=13, color=WHITE, bold=True)
# -- left column: structured data
text(s, X0, 6.7, 4.3, 0.4, 'Structured data', size=16, color=BLUE, bold=True, align='c')
rect(s, X0, 7.1, 4.3, 1.95, fill=None, radius=0.25, line=BLUE, line_w=1.25, dash=MSO_LINE.ROUND_DOT)
d2 = rect(s, X0 + 0.3, 7.3, 3.7, 0.7, fill=PANEL, radius=0.12)
shape_text(d2, 'National HIE  ·  diabetes registry', size=13, color=INK, bold=True)
text(s, X0 + 0.3, 8.05, 3.7, 0.75, 'Structured tabular data: patients, encounters, diagnoses, facilities', size=11, color=SLATE, align='c')
# -- agentic retrieval arrow + agent
AX, AY, AW, AH = 7.0, 3.0, 4.4, 5.35
text(s, X0 + 4.3, 4.45, AX - X0 - 4.3, 0.4, 'Agentic retrieval', size=12, color=BLUE, bold=True, align='c')
arrow(s, X0 + 4.3 + 0.08, 4.85, AX - X0 - 4.3 - 0.16, 0.36, fill=BLUE)
card(s, AX, AY, AW, AH, fill=BLUE)
text(s, AX + 0.3, AY + 0.25, AW - 0.6, 0.55, 'Population Health Agent', size=20, color=WHITE, bold=True, align='c')
rect(s, AX + 0.9, AY + 0.85, AW - 1.8, 0.03, fill=WHITE, radius=0)
text(s, AX + 0.3, AY + 1.0, AW - 0.6, 1.45, 'Patient-level queries, population insights, cost forecasting and policy simulation, with traceable, cited reasoning', size=14, color=WHITE, align='c', anchor='t')
ic_ = rect(s, AX + 0.3, AY + 2.6, AW - 0.6, 2.5, fill=PALE, radius=0.18)
text(s, AX + 0.5, AY + 2.75, AW - 1.0, 2.25, [
    {'t': '**Orchestrator**  ·  picks the tool for each step', 'bullet': True},
    {'t': '**Guideline grounding**  ·  cites document and section', 'bullet': True, 'space_before': 4},
    {'t': '**Audit log**  ·  every call visible in the trace', 'bullet': True, 'space_before': 4},
], size=13, color=NAVY, anchor='m')
# -- MCP server
MX, MY, MW, MH = 12.3, 3.85, 2.3, 2.3
arrow(s, AX + AW + 0.08, 4.85, MX - AX - AW - 0.16, 0.36, fill=BLUE)
m = rect(s, MX, MY, MW, MH, fill=LIGHT, radius=0.3, shadow=True)
icon(s, 'usb', 'n', MX + MW / 2 - 0.4, MY + 0.3, 0.8)
text(s, MX, MY + 1.2, MW, 0.9, ['MCP server', {'t': 'one plug, every tool', 'size': 11, 'bold': False}], size=18, color=NAVY, bold=True, align='c', anchor='t')
# -- tools
TX, TW, TH, TG = 15.4, X1 - 15.4, 0.72, 0.22
text(s, TX, 2.9, TW, 0.4, 'Tools', size=16, color=BLUE, bold=True, align='c')
tools5 = [('Run & build ML model (forecasting, simulation)', SLATE, WHITE),
          ('Run decision flow (SAS Intelligent Decisioning)', LIGHT, NAVY),
          ('Generate charts', SKY, WHITE),
          ('Query data (SQL)', BLUE, WHITE)]
tool_y = []
for i, (t, col, tc) in enumerate(tools5):
    yy = 3.4 + i * (TH + TG); tool_y.append(yy)
    tt = rect(s, TX, yy, TW, TH, fill=col, radius=0.1, shadow=True)
    shape_text(tt, t, size=12, color=tc, bold=True, margins=(0.1, 0.03, 0.1, 0.03))
    connector(s, MX + MW + 0.04, MY + MH / 2, TX - 0.05, yy + TH / 2, color=BLUE, width=1.5, tail=True)
# -- CAS table under the agent, fed by the structured data, read by the SQL tool
icon(s, 'table', 'b', AX + 0.35, 8.52, 0.5)
text(s, AX + 0.95, 8.5, 3.6, 0.5, 'CAS table (SAS Viya)', size=16, color=BLUE, bold=True, anchor='m')
connector(s, X0 + 4.3, 8.0, AX + 0.3, 8.75, color=BLUE, width=1.5, dash=MSO_LINE.ROUND_DOT, tail=True)
connector(s, AX + 3.9, 8.75, TX + TW / 2, tool_y[3] + TH + 0.02, color=BLUE, width=1.5, dash=MSO_LINE.ROUND_DOT, tail=True)
# -- what it delivers
callout(s, X0, 9.15, XW, 1.15, [
    {'t': 'What this solution delivers', 'bold': True, 'size': 15},
    {'t': 'Natural-language questions → governed analytics, cohorts and charts   ·   ML scoring, forecasting and policy simulation on demand   ·   '
          'Every answer cites the national guideline and shows its tool calls', 'size': 13}], icon_name='badge', size=14)
notes(s, """
The one picture to remember, and the one to copy for your own use case. Left: two kinds of data. The guidelines are
chunked, embedded and stored as vectors (the RAG knowledge base); the registry is loaded as a CAS table in SAS Viya.
Middle: the agent, with its orchestrator, its guideline grounding and its audit log. Right: the MCP server and the
tools it exposes: query the data, score the model, run a decision flow, draw a chart. Follow the dotted line: the
registry becomes a CAS table, and the SQL tool reads it; the agent itself never touches the table. The strip at the
bottom is what the solution delivers. Everything on this slide is what the participants assemble in the hands-on.
""")

# 16 ---- Demo plan --------------------------------------------------------------------------
s = content_slide('What you will see in the demo', 'Five real questions, and what to watch for in the agent\'s trace')
rows_ = [('#', 'The question', 'What the agent does', 'Watch for'),
         ('1', 'How many patients are in the registry, and how many are well controlled?', 'One SQL query on the registry; cites NHA-CG-01 §2 for the HbA1c bands', 'The query in the trace; the cohort stated next to every number'),
         ('2', 'Compare glycaemic control across the emirates.', 'A group-by query, a chart, and the equity check against NHA-PP-01 §4', 'The 5-point equity trigger, named without being asked'),
         ('3', 'What is the LDL target for a very-high-risk patient?', 'Retrieval only: no query; the answer comes from NHA-CG-02 §3', 'The citation, and “not covered” when the collection is silent'),
         ('4', 'Who would qualify for a GLP-1 RA under the guideline, and how many are already on one?', 'Guideline first (NHA-CG-01 §4), then two queries: eligible, and not yet treated', 'RAG and SQL in one answer; the clinician decides (§7)'),
         ('5', 'What is patient EHS-100092\'s risk of deterioration in the next 12 months?', 'Scores the patient through the published model and contrasts it with the registry\'s rule-based tier', 'The model call in the trace; restricted-consent patients are never listed')]
tx, ty, tw_, th = X0, 3.0, XW, 6.1
tbl = s.shapes.add_table(len(rows_), 4, Inches(tx), Inches(ty), Inches(tw_), Inches(th)).table
colw = [0.75, 6.3, 5.6, 4.6]
for i, w_ in enumerate(colw): tbl.columns[i].width = Inches(w_)
tbl.rows[0].height = Inches(0.6)
for r in range(1, len(rows_)): tbl.rows[r].height = Inches((th - 0.6) / (len(rows_) - 1))
for r, row in enumerate(rows_):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (WHITE if r % 2 == 1 else ARC))
        cell.margin_left = Inches(0.18); cell.margin_right = Inches(0.18); cell.margin_top = Inches(0.06); cell.margin_bottom = Inches(0.06)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True
        if r == 0:
            fill_paras(tf, val, size=15, color=WHITE, bold=True, align='c' if c == 0 else 'l')
        else:
            fill_paras(tf, val, size=13, color=(BLUE if c == 0 else (NAVY if c == 1 else INK)), bold=(c in (0, 1)), align='c' if c == 0 else 'l')
tblPr = tbl._tbl.tblPr
tblPr.set('firstRow', '1'); tblPr.set('bandRow', '0')
text(s, X0, 9.3, XW, 0.9, [{'t': 'If time allows, question 6: “If we put every eligible uncontrolled patient on a GLP-1 RA and closed the statin gap, what would it save?” '
     'A policy what-if: the agent takes the measured effects from NHA-PP-01 §5, scales them to the cohorts it counted, and shows the arithmetic.', 'size': 13, 'color': SLATE}])
notes(s, """
One minute on this slide, then switch to SAS Retrieval Agent Manager. Five questions in a deliberate order: a plain
count (so the room sees a query and a cohort), a comparison (a chart and the equity rule firing on its own), a pure
guideline question (no query at all, just the citation), a question that needs RAG and SQL together, and a model
score for one patient. For each question, show the answer first, then open the trace and point at the tool call or
the retrieved passage. Expected answers are in the kit (agents/population_health_agent.md): 4,000 patients, 34.9% well
controlled; Ajman 46.3% to Umm Al Quwain 22.0%; LDL below 1.4 mmol/L; 399 eligible, 373 not yet treated; EHS-100092
well above the 10.4% base rate. If the environment is slow, drop question 2; if it fails, use slide 17.
""")

# 17 ---- Under the hood ---------------------------------------------------------------------
s = content_slide('Under the hood of one answer', 'Question 4 step by step, from retrieval to the hand-over')
rect(s, X0, 2.85, XW, 7.45, fill=PANEL, radius=0.35, shadow=True)
hsteps = [('1 · Retrieve', 'Searches the NHA collection. NHA-CG-01 §4: consider a GLP-1 RA when HbA1c stays above target and BMI is over 30.', BLUE),
          ('2 · Query', 'query_data: type 2, HbA1c ≥ 9%, BMI > 30, from Public.EHS_DIABETES. Result: 399 patients.', NAVY),
          ('3 · Query again', 'Same cohort, on_sglt2_glp1 = 0. Result: 373 not yet treated. Restricted-consent patients are counted, never listed.', NAVY),
          ('4 · Score', 'score_table_rows on the published model for the ten highest-cost eligible patients: a probability and its drivers for each.', SKY),
          ('5 · Write', 'The answer: the cohort in words, the two numbers with their denominator, the citation, and the source line (query + sections).', BLUE),
          ('6 · Hand over', 'A draft review list for the clinic. The guideline (§7) and the agent\'s own rules say: a clinician decides; the agent never prescribes.', GREEN)]
cw2 = 5.05; gap = 0.67; ch2 = 2.75
for i, (t, d, col) in enumerate(hsteps):
    r, c = divmod(i, 3)
    x = X0 + 0.5 + c * (cw2 + gap); y = 3.65 + r * (ch2 + 0.95)
    card(s, x, y, cw2, ch2, shadow=False)
    p = rect(s, x + 0.35, y - 0.3, cw2 - 0.7, 0.6, fill=col, radius=0.3, shadow=True)
    shape_text(p, t, size=15, color=WHITE, bold=True, margins=(0.1, 0.02, 0.1, 0.02))
    text(s, x + 0.35, y + 0.5, cw2 - 0.7, ch2 - 0.65, d, size=14, color=INK, anchor='m', align='c')
    if c < 2:
        arrow(s, x + cw2 + 0.12, y + ch2 / 2 - 0.22, gap - 0.24, 0.44, fill=BLUE)
notes(s, """
Use this slide in one of two ways. After a successful demo: replay question 4 as six beats and point out that this is
the loop from slide 5, two full turns plus the close. If the live environment failed: this is the demo, told as a
story, with the real numbers. Three design decisions hide in it: the guideline was retrieved before any number was
produced (grounding), restricted-consent patients were counted but never listed (a rule the tool enforces), and the
output is a draft for a clinician (human oversight). Ask: where could this have gone wrong without those three?
""")

# 18 ---- Guardrails -------------------------------------------------------------------------
s = content_slide('The guardrails you saw in the demo', 'Written into the agent\'s instructions and enforced by the tools')
guards = [('calculator', 'Numbers come from queries', 'The agent runs query_data before stating any count, rate or cost. No estimates, no reuse of an old number for a new cohort.', 'Lives in: the agent instructions, rule 1'),
          ('cite', 'Clinical statements are cited', 'Document and section, or “not covered by the NHA guidelines in my collection”. Never from general knowledge.', 'Lives in: the agent instructions, rule 3'),
          ('lock', 'Consent is respected', 'Restricted-consent patients appear in aggregates only: never named, listed or scored (NHA-PP-01 §6).', 'Lives in: the instructions, and the query the tool runs'),
          ('users', 'Humans decide', 'The agent drafts recall lists and recommendations; a clinician or the programme lead approves them (NHA-CG-01 §7).', 'Lives in: the agent instructions, rule 5'),
          ('key', 'Tools are scoped', 'The MCP server sees only the team\'s registry table and model. Anything else answers “out of scope”, by design.', 'Lives in: the MCP server (allowed tables and models)'),
          ('log', 'The trace is visible', 'Every retrieval, query and model call is shown next to the answer. That is the audit trail, built in, not bolted on.', 'Lives in: SAS Retrieval Agent Manager, on every answer')]
gw = (XW - 2 * 0.4) / 3; gh = 3.3
for i, (ic, t, d, where) in enumerate(guards):
    r, c = divmod(i, 3)
    x = X0 + c * (gw + 0.4); y = 2.95 + r * (gh + 0.4)
    card(s, x, y, gw, gh)
    icon_circle(s, x + 0.65, y + 0.62, 0.72, [NAVY, BLUE, SKY][c], ic, shadow=False, scale=0.55)
    text(s, x + 1.2, y + 0.35, gw - 1.45, 0.55, t, size=18, color=INK, bold=True)
    text(s, x + 0.4, y + 1.2, gw - 0.8, 1.4, d, size=14, color=INK)
    text(s, x + 0.4, y + gh - 0.75, gw - 0.8, 0.5, [{'t': where, 'color': BLUE, 'size': 13, 'bold': True}], anchor='b')
notes(s, """
Six guardrails, and the point is where they live: in the agent's written instructions (the system prompt the
participants will paste) and in the tools (the MCP server's scope, the consent filter), not in the model's good
judgement. Numbers from queries, citations or silence, consent, humans decide, scoped tools, a visible trace. These
six are also the answer to 'how do we audit an agent?': the trace shows every call, the scope shows what it could
reach, and the instructions show what it was told. Carry this list into the hands-on: each team writes these rules
into its own agent.
""")

# 19 ---- Section: what you will build -------------------------------------------------------
s = section_slide('What you will build', 'The same agent, assembled by you, without code', 'rocket')
notes(s, "Part 4 of 4. About 4 minutes: the five build steps and the component table, then the takeaways.")

# 20 ---- Build path -------------------------------------------------------------------------
s = content_slide('Your build path: five steps, no code', 'From a CSV and four PDFs to a working agent in SAS Viya and SAS RAM')
steps5 = [('upload', 'Load the data', 'Import the registry CSV into SAS Viya (Manage Data). It becomes a CAS table that the tools can query.', 'SAS Viya  ·  structured data', NAVY),
          ('model', 'Build and publish the model', 'Model Studio trains the pipeline; the champion predicts deterioration in 12 months; publish it for scoring.', 'SAS Viya  ·  machine learning', BLUE),
          ('files', 'Index the guidelines', 'Create a RAM collection from the four NHA PDFs: chunk, embed, store. Write the retrieval instructions.', 'SAS RAM  ·  RAG', SKY),
          ('api', 'Connect the tools', 'Register the MCP tool server, scoped to your table and your model: list, describe, query, score.', 'SAS RAM  ·  MCP', BLUE),
          ('robot', 'Assemble and test the agent', 'Instructions + collection + tools. Run the test questions and read the trace behind each answer.', 'SAS RAM  ·  agent', NAVY)]
sw5 = (XW - 4 * 0.42) / 5
for i, (ic, t, d, tag, col) in enumerate(steps5):
    x = X0 + i * (sw5 + 0.42); y = 3.55; h = 5.3
    card(s, x, y, sw5, h)
    icon_circle(s, x + sw5 / 2, y, 1.05, col, ic)
    c = oval(s, x + 0.45, y + 0.45, 0.5, fill=PALE); shape_text(c, str(i + 1), size=14, color=NAVY, bold=True, margins=(0, 0, 0, 0))
    text(s, x + 0.25, y + 0.85, sw5 - 0.5, 0.95, t, size=17, color=INK, bold=True, align='c')
    text(s, x + 0.3, y + 1.85, sw5 - 0.6, 2.5, d, size=13, color=INK)
    p = rect(s, x + 0.3, y + h - 0.85, sw5 - 0.6, 0.5, fill=LIGHT, radius=0.25)
    shape_text(p, tag, size=11, color=NAVY, bold=True, margins=(0.05, 0.02, 0.05, 0.02))
    if i < 4:
        chevron(s, x + sw5 + 0.06, y + h / 2 - 0.22, 0.3, 0.44, fill=BLUE)
callout(s, X0, 9.2, XW, 1.0, 'A ready-made copilot agent in RAM can drive steps 1 and 2 for you through MCP: talk to your data, build the model. Steps 3 to 5 are yours, in the RAM interface.', icon_name='wand', size=15)
notes(s, """
Five steps, in the order the participants will do them. Steps 1 and 2 live in SAS Viya: load the registry, build and
publish the model. Steps 3 to 5 live in SAS Retrieval Agent Manager: the collection (RAG), the MCP tool source, and
the agent itself. Point at the coloured tags: the colours match the building blocks from part 2. Mention the copilot:
a pre-built agent in RAM can do the data loading and the model build through MCP on request, which is itself a nice
demonstration of an agent at work. The hands-on guide has a screenshot for every click.
""")

# 21 ---- Component table --------------------------------------------------------------------
s = content_slide('The agent on one page', 'What each part gives the agent, where you build it, and the material to use')
rows_ = [('Component', 'What it gives the agent', 'Where you build it', 'Material in the kit'),
         ('Structured data', 'Every number: cohorts, rates, costs, the patient a clinician asks about', 'SAS Viya  ·  CAS table (Public.EHS_DIABETES)', 'ehs_diabetes_registry.csv (4,000 rows, 54 columns), ehs_facilities.csv, the data dictionary'),
         ('ML model', 'The 12-month deterioration risk of a patient or a cohort, with its drivers', 'SAS Viya  ·  Model Studio, champion published to SAS Micro Analytic Service', 'Target deterioration_next_12m; leakage columns listed in the data dictionary'),
         ('Knowledge (RAG)', 'Guideline answers with document and section, or “not covered”', 'SAS RAM  ·  collection + retrieval settings', 'NHA-CG-01, NHA-CG-02, NHA-CG-03, NHA-PP-01 as PDFs'),
         ('Tools (MCP)', 'Its hands: list, describe and preview tables, query, list and describe models, score', 'SAS RAM  ·  MCP tool source, scoped to your table and model', 'The Bootcamp MCP server (8 tools), or the SAS Viya MCP server'),
         ('The agent', 'The orchestrator: reads the question, picks the tool, writes the cited answer', 'SAS RAM  ·  agent = instructions + collection + tools', 'The system prompt template and 20 test questions with expected answers')]
tx, ty, tw_, th = X0, 3.0, XW, 6.9
tbl = s.shapes.add_table(len(rows_), 4, Inches(tx), Inches(ty), Inches(tw_), Inches(th)).table
colw = [2.9, 5.0, 4.75, 4.6]
for i, w_ in enumerate(colw): tbl.columns[i].width = Inches(w_)
tbl.rows[0].height = Inches(0.65)
for r in range(1, len(rows_)): tbl.rows[r].height = Inches((th - 0.65) / (len(rows_) - 1))
for r, row in enumerate(rows_):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (WHITE if r % 2 == 1 else ARC))
        cell.margin_left = Inches(0.2); cell.margin_right = Inches(0.2); cell.margin_top = Inches(0.08); cell.margin_bottom = Inches(0.08)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True
        if r == 0:
            fill_paras(tf, val, size=15, color=WHITE, bold=True)
        else:
            fill_paras(tf, val, size=13, color=(BLUE if c == 2 else INK), bold=(c == 0))
tblPr = tbl._tbl.tblPr
tblPr.set('firstRow', '1'); tblPr.set('bandRow', '0')
notes(s, """
The reference slide for the hands-on: one row per component, what it gives the agent, where it is built, and the
material in the kit. Skippable if you are behind; the hands-on guide repeats all of it with screenshots. If you keep
it, read only the third column: everything in SAS Viya is data and model, everything in RAM is knowledge, tools and
the agent.
""")

# 22 ---- Takeaways --------------------------------------------------------------------------
s = prs.slides.add_slide(LAY['SAS - Blank - White'])
rect(s, 0, 0, 20, 11.25, fill=BLUE, radius=0)
rect(s, 14.2, 0, 5.8, 11.25, fill=NAVY, radius=0)
text(s, 1.37, 0.85, 12.0, 1.0, 'Takeaways', size=56, color=WHITE, bold=True, anchor='m')
tks = [('01', 'An agent is a loop with a goal, tools and boundaries', 'It reasons, acts, checks, and keeps going until the job is done or a person must step in.'),
       ('02', 'The LLM reasons; RAG, tools and models keep it honest', 'Words from the LLM, facts from your documents, numbers from queries and models, all connected through MCP and logged.'),
       ('03', 'You will build this agent in the bootcamp', 'Load the data, publish the model, index the guidelines, connect the tools, assemble and test.')]
for i, (n, a, b) in enumerate(tks):
    y = 2.9 + i * 2.35
    text(s, 1.37, y - 0.15, 2.2, 1.4, n, size=60, color=LIGHT, bold=True, anchor='t')
    text(s, 3.7, y, 9.9, 0.7, a, size=24, color=WHITE, bold=True)
    text(s, 3.7, y + 0.7, 9.9, 1.2, b, size=17, color=LIGHT)
icon(s, 'arrowRight', 'w', 15.2, 3.6, 0.9)
text(s, 15.2, 4.7, 4.2, 0.6, 'Up next  ·  hands-on', size=16, color=LIGHT, bold=True)
text(s, 15.2, 5.3, 4.2, 2.2, 'Build your own Population Health Agent', size=26, color=WHITE, bold=True)
text(s, 15.2, 7.4, 4.2, 1.6, 'Bring the registry CSV, the four NHA guideline PDFs and the test questions.', size=15, color=LIGHT)
copyright_line(s, color=LIGHT)
notes(s, """
Three sentences to remember. If they remember only one: the LLM reasons, and everything around it keeps it honest.
Then hand over to the hands-on part: the participants build exactly the agent they have just seen, step by step, with
the materials in the kit.
""")

# 23 ---- Closing ----------------------------------------------------------------------------
s = prs.slides.add_slide(LAY['SAS - Closing'])
set_ph(s.shapes.title, 'Thank you!', size=72)
set_ph(placeholder(s, 11), 'Questions?', size=28)
notes(s, "Open questions for two or three minutes, then move to the hands-on build.")

prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides))
