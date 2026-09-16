#!/usr/bin/env python3
"""Build 'Session 2: How an AI agent works and makes decisions' on the SAS EXTERNAL template.

Design language copied from 'SAS Viya Agentic AI Experience - Presentation':
  * SAS 2023 palette (Blue 0766D1, Midnight 032954, Sky 4398F9, Light C4DEFD, Slate 7E889A)
  * Anova Bold / Anova Light theme fonts (embedded in the template)
  * Slate 56pt title + blue 36pt subtitle, light-grey arc band, white rounded cards with soft
    shadows, icon circles overlapping card tops, navy pill headers with blue chevrons,
    slate section dividers, blue title/closing/takeaway slides.
"""
import json, os, re, math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ICONS = os.path.join(HERE, 'icons')                       # PNG icons rendered by icons.js
ARC_JSON = os.path.join(HERE, 'arc.json')                  # light-grey arc band traced from the reference deck
TEMPLATE = os.environ.get('SAS_TEMPLATE', os.path.normpath(os.path.join(HERE, '..', '..', '..', 'templates', 'SAS_External_Template.pptx')))
OUT = os.environ.get('SESSION2_OUT', os.path.normpath(os.path.join(HERE, '..', 'Session2_How_an_AI_Agent_Works_and_Makes_Decisions.pptx')))

# ---- palette (SAS-2023 theme) --------------------------------------------------------------
BLUE = '0766D1'; NAVY = '032954'; SKY = '4398F9'; LIGHT = 'C4DEFD'; MID = '98C5FB'
SLATE = '7E889A'; SLATE2 = '838E9E'; GRAY = 'BAC0C9'; PANEL = 'E4E7EA'; ARC = 'F0F1F3'
WHITE = 'FFFFFF'; BLACK = '000000'; INK = '262626'; GREEN = '009242'; RED = 'C00000'
PALE = 'EAF3FE'

# ---- page geometry (20 x 11.25 in) ---------------------------------------------------------
X0 = 1.37; XW = 17.25; X1 = X0 + XW; CT = 2.85; CB = 10.4
MJ = '+mj-lt'   # theme major font (Anova Bold)

def rgb(h): return RGBColor.from_string(h)

# ============================================================================================
# low-level helpers
# ============================================================================================
def _fill(shape, hexcol):
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(hexcol)

def _no_line(shape):
    shape.line.fill.background()

def _line(shape, hexcol, width=1.0, dash=None):
    shape.line.color.rgb = rgb(hexcol); shape.line.width = Pt(width)
    if dash: shape.line.dash_style = dash

def _effect_clear(spPr):
    for el in spPr.findall(qn('a:effectLst')):
        spPr.remove(el)

def _no_shadow(shape):
    spPr = shape._element.spPr
    _effect_clear(spPr)
    etree.SubElement(spPr, qn('a:effectLst'))

def _shadow(shape, blur=16, dist=3, alpha=16, angle=90):
    spPr = shape._element.spPr
    _effect_clear(spPr)
    eff = etree.SubElement(spPr, qn('a:effectLst'))
    sh = etree.SubElement(eff, qn('a:outerShdw'))
    sh.set('blurRad', str(int(Pt(blur)))); sh.set('dist', str(int(Pt(dist))))
    sh.set('dir', str(angle * 60000)); sh.set('algn', 'ctr'); sh.set('rotWithShape', '0')
    clr = etree.SubElement(sh, qn('a:srgbClr')); clr.set('val', '000000')
    a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(alpha * 1000))

def _bullet(p, color=BLUE, indent=0.26, char='•'):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(Inches(indent)))); pPr.set('indent', str(-int(Inches(indent))))
    for tag in ('a:buClr', 'a:buFont', 'a:buChar', 'a:buNone'):
        for el in pPr.findall(qn(tag)): pPr.remove(el)
    buClr = etree.SubElement(pPr, qn('a:buClr')); c = etree.SubElement(buClr, qn('a:srgbClr')); c.set('val', color)
    buFont = etree.SubElement(pPr, qn('a:buFont')); buFont.set('typeface', 'Arial')
    buChar = etree.SubElement(pPr, qn('a:buChar')); buChar.set('char', char)

_ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
_ANCHOR = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}

def _runs(p, text, size, color, bold, font, italic=False):
    """'**bold**' markup -> Anova Bold runs."""
    parts = text.split('**')
    for i, seg in enumerate(parts):
        if seg == '': continue
        r = p.add_run(); r.text = seg
        f = r.font; f.size = Pt(size); f.color.rgb = rgb(color)
        is_b = bold or (i % 2 == 1)
        if is_b: f.name = MJ
        elif font: f.name = font
        if italic: f.italic = True

def fill_paras(tf, paras, size=18, color=BLACK, bold=False, font=None, align='l',
               line_spacing=None, space_after=None, space_before=None):
    if isinstance(paras, str): paras = [paras]
    first = True
    for item in paras:
        if isinstance(item, str): item = {'t': item}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = _ALIGN[item.get('align', align)]
        ls = item.get('line_spacing', line_spacing)
        if ls: p.line_spacing = ls
        sa = item.get('space_after', space_after)
        if sa is not None: p.space_after = Pt(sa)
        sb = item.get('space_before', space_before)
        if sb is not None: p.space_before = Pt(sb)
        if item.get('bullet'):
            _bullet(p, item.get('bullet_color', BLUE), indent=item.get('indent', 0.26), char=item.get('char', '•'))
        _runs(p, item.get('t', ''), item.get('size', size), item.get('color', color),
              item.get('bold', bold), item.get('font', font), item.get('italic', False))
    return tf

def text(slide, x, y, w, h, paras, size=18, color=BLACK, bold=False, font=None, align='l', anchor='t',
         margins=(0.05, 0.03, 0.05, 0.03), line_spacing=None, space_after=None, rotation=None, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = _ANCHOR[anchor]
    fill_paras(tf, paras, size, color, bold, font, align, line_spacing, space_after)
    if rotation is not None: tb.rotation = rotation
    return tb

def shape_text(shape, paras, size=18, color=WHITE, bold=False, font=None, align='c', anchor='m',
               margins=(0.12, 0.06, 0.12, 0.06), line_spacing=None, space_after=None):
    tf = shape.text_frame; tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = _ANCHOR[anchor]
    fill_paras(tf, paras, size, color, bold, font, align, line_spacing, space_after)
    return shape

def rect(slide, x, y, w, h, fill=WHITE, radius=0.28, shadow=False, line=None, line_w=1.0, dash=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius: s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill: _fill(s, fill)
    else: s.fill.background()
    if line: _line(s, line, line_w, dash)
    else: _no_line(s)
    if shadow: _shadow(s)
    else: _no_shadow(s)
    return s

def card(slide, x, y, w, h, fill=WHITE, radius=0.28, shadow=True, line=None):
    return rect(slide, x, y, w, h, fill=fill, radius=radius, shadow=shadow, line=line)

def pill(slide, x, y, w, h, label, fill=NAVY, color=WHITE, size=16, bold=True, shadow=False):
    s = rect(slide, x, y, w, h, fill=fill, radius=h / 2, shadow=shadow)
    shape_text(s, label, size=size, color=color, bold=bold, margins=(0.1, 0.02, 0.1, 0.02))
    return s

def oval(slide, cx, cy, d, fill=BLUE, shadow=False, line=None, line_w=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    if fill: _fill(s, fill)
    else: s.fill.background()
    if line: _line(s, line, line_w)
    else: _no_line(s)
    if shadow: _shadow(s)
    else: _no_shadow(s)
    return s

def icon(slide, name, variant, x, y, size):
    path = os.path.join(ICONS, f'{name}_{variant}.png')
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(size), Inches(size))

def icon_circle(slide, cx, cy, d, fill, name, variant='w', shadow=True, scale=0.52):
    oval(slide, cx, cy, d, fill=fill, shadow=shadow)
    isz = d * scale
    icon(slide, name, variant, cx - isz / 2, cy - isz / 2, isz)

def arrow(slide, x, y, w, h, fill=BLUE, rotation=0):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(s, fill); _no_line(s); _no_shadow(s)
    if rotation: s.rotation = rotation
    return s

def chevron(slide, x, y, w, h, fill=BLUE, rotation=0):
    s = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(s, fill); _no_line(s); _no_shadow(s)
    if rotation: s.rotation = rotation
    return s

def connector(slide, x1, y1, x2, y2, color=SLATE, width=1.25, dash=None, head=False, tail=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    if dash: c.line.dash_style = dash
    ln = c.line._get_or_add_ln()
    if head:
        e = etree.SubElement(ln, qn('a:headEnd')); e.set('type', 'triangle'); e.set('w', 'med'); e.set('len', 'med')
    if tail:
        e = etree.SubElement(ln, qn('a:tailEnd')); e.set('type', 'triangle'); e.set('w', 'med'); e.set('len', 'med')
    spPr = c._element.spPr
    _effect_clear(spPr); etree.SubElement(spPr, qn('a:effectLst'))
    return c

def send_to_back(slide, shape, index=2):
    sp = shape._element; tree = slide.shapes._spTree
    tree.remove(sp); tree.insert(index, sp)

ARCPTS = json.load(open(ARC_JSON))['points']
def arc_bg(slide):
    scale = 12700  # EMU per PDF point; reference page = 1440 x 810 pt = 20 x 11.25 in
    fb = slide.shapes.build_freeform(int(ARCPTS[0][0] * scale), int(ARCPTS[0][1] * scale), scale=1.0)
    fb.add_line_segments([(int(x * scale), int(y * scale)) for x, y in ARCPTS[1:]], close=True)
    s = fb.convert_to_shape()
    _fill(s, ARC); _no_line(s); _no_shadow(s)
    send_to_back(slide, s)
    return s

def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt.strip()

def placeholder(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx: return ph
    return None

def drop_placeholder(ph):
    if ph is not None:
        ph._element.getparent().remove(ph._element)

def set_ph(ph, txt, size=None, color=None, bold=None):
    tf = ph.text_frame
    tf.text = txt
    for p in tf.paragraphs:
        for r in p.runs:
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = rgb(color)
            if bold is not None: r.font.bold = bold

# ============================================================================================
# presentation set-up
# ============================================================================================
prs = Presentation(TEMPLATE)
# drop template slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
# keep only master 0 (SAS - EXTERNAL)
mlst = prs.slide_masters._sldMasterIdLst
for el in list(mlst)[1:]:
    prs.part.drop_rel(el.rId); mlst.remove(el)
M = prs.slide_masters[0]
LAY = {l.name: l for l in M.slide_layouts}

def content_slide(title, subtitle=None, arc=True, title_size=None, subtitle_size=30):
    s = prs.slides.add_slide(LAY['SAS - Title & Subtitle'])
    if arc: arc_bg(s)
    t = s.shapes.title
    t.text = title
    if not title_size and len(title) > 40:
        title_size = 48   # keep long titles on one line (56pt fits ~40 chars at 17.25in)
    if title_size:
        for p in t.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(title_size)
    ph = placeholder(s, 10)
    if subtitle:
        set_ph(ph, subtitle, size=subtitle_size)
    else:
        drop_placeholder(ph)
    return s

def section_slide(title, subtitle, icon_name, title_size=60):
    s = prs.slides.add_slide(LAY['SAS - Section'])
    set_ph(s.shapes.title, title, size=title_size)
    set_ph(placeholder(s, 10), subtitle, size=36)
    # outlined icon top-right like the reference dividers
    oval(s, 17.15, 2.15, 2.3, fill=None, line=WHITE, line_w=2.5)
    icon(s, icon_name, 'w', 17.15 - 0.62, 2.15 - 0.62, 1.24)
    return s

def heading(slide, x, y, w, txt, size=24, color=INK, h=0.5, align='l'):
    return text(slide, x, y, w, h, txt, size=size, color=color, bold=True, align=align, anchor='t')

def body(slide, x, y, w, h, paras, size=17, color=INK, space_after=6, anchor='t'):
    return text(slide, x, y, w, h, paras, size=size, color=color, space_after=space_after, anchor=anchor)

def icon_row(slide, x, y, w, name, label, desc, circle=BLUE, d=0.7, label_size=18, desc_size=15,
             variant='w', row_h=None, desc_color=INK):
    icon_circle(slide, x + d / 2, y + d / 2, d, circle, name, variant=variant, shadow=False, scale=0.55)
    text(slide, x + d + 0.22, y - 0.06, w - d - 0.22, 0.42, label, size=label_size, color=INK, bold=True)
    tb = text(slide, x + d + 0.22, y + 0.34, w - d - 0.22, (row_h or 1.2) - 0.34, desc, size=desc_size, color=desc_color)
    return tb

def numbered_row(slide, x, y, w, n, label, desc, d=0.62, label_size=18, desc_size=15, row_h=1.0, fill=BLUE):
    c = oval(slide, x + d / 2, y + d / 2, d, fill=fill)
    shape_text(c, str(n), size=18, color=WHITE, bold=True, margins=(0, 0, 0, 0))
    text(slide, x + d + 0.22, y - 0.05, w - d - 0.22, 0.42, label, size=label_size, color=INK, bold=True)
    if desc:
        text(slide, x + d + 0.22, y + 0.34, w - d - 0.22, row_h - 0.34, desc, size=desc_size, color=INK)

def callout(slide, x, y, w, h, paras, fill=LIGHT, color=NAVY, size=18, icon_name=None, icon_variant='n', align='l'):
    r = rect(slide, x, y, w, h, fill=fill, radius=0.22)
    pad = 0.35
    if icon_name:
        isz = min(0.7, h - 0.4)
        icon(slide, icon_name, icon_variant, x + 0.35, y + (h - isz) / 2, isz)
        pad = 0.35 + isz + 0.3
    text(slide, x + pad, y + 0.12, w - pad - 0.3, h - 0.24, paras, size=size, color=color, anchor='m', align=align)
    return r

def copyright_line(slide, color=LIGHT):
    text(slide, 1.37, 10.78, 6.0, 0.27, 'Copyright © SAS Institute Inc. All rights reserved.', size=10, color=color, anchor='b')

# ============================================================================================
# SLIDES
# ============================================================================================
# 1 ---- Title -------------------------------------------------------------------------------
s = prs.slides.add_slide(LAY['SAS - Title'])
set_ph(s.shapes.title, 'How an AI agent works\nand makes decisions', size=66)
set_ph(placeholder(s, 10), 'Session 2  ·  Agentic AI Bootcamp  ·  Day 1', size=40)
set_ph(placeholder(s, 11), 'Raed Aldweik  |  SAS\nEHS Agentic AI Bootcamp', size=24)
notes(s, """
Welcome back after the coffee break. Session 1 gave you the big picture: what Agentic AI is, how it differs from
generative AI and automation, and the AI lifecycle. In the next 45 minutes we open the agent up and look inside.
Promise for the session: by the end you will be able to describe any agent using six building blocks, follow the loop it
runs, and decide how much freedom to give it and when a person must step in. No code, no maths.
Timing guide: intro 3 min, anatomy 14 min, decision loop 9 min, enterprise integration 5 min, human-in-the-loop 7 min,
boundaries & escalation 7 min, exercise 5 min, wrap-up 2 min (~52 min). For a strict 45 minutes, skip slides 3 and 21
and shorten the exercise to 3 minutes.
""")

# 2 ---- Agenda ------------------------------------------------------------------------------
s = content_slide('In the next 45 minutes', 'Five questions, one running example, and a short exercise')
items = [
    ('Anatomy of an AI agent', 'Goal, Planning, Knowledge, Tools, Actions and Human Oversight'),
    ('How an agent decides and executes', 'The think, act, observe loop and what happens inside one step'),
    ('Connecting to enterprise systems, data and tools', 'How the agent reaches your systems safely'),
    ('Human-in-the-loop vs autonomous execution', 'The autonomy spectrum and how to choose a level'),
    ('Autonomy, boundaries and escalation points', 'The fence around the agent, and the door out'),
]
connector(s, 1.82, 3.35, 1.82, 8.95, color=LIGHT, width=2.5)
for i, (a, b) in enumerate(items):
    y = 3.0 + i * 1.42
    sq = rect(s, 1.37, y, 0.9, 0.9, fill=BLUE, radius=0.12)
    shape_text(sq, str(i + 1), size=26, color=WHITE, bold=True, margins=(0, 0, 0, 0))
    text(s, 2.55, y - 0.04, 8.2, 0.5, a, size=22, color=INK, bold=True)
    text(s, 2.55, y + 0.42, 8.2, 0.5, b, size=15, color=SLATE)
card(s, 11.4, 2.95, 7.22, 7.0)
heading(s, 11.85, 3.3, 6.4, 'By the end of this session you can…', size=22, color=BLUE)
outs = [
    ('Explain what is inside an agent', 'using the six building blocks and one healthcare example'),
    ('Read an agent\'s decision loop', 'and spot the points where it can go wrong'),
    ('Decide how much autonomy to give', 'an agent for each action, and where a person must step in'),
]
for i, (a, b) in enumerate(outs):
    y = 4.25 + i * 1.75
    icon_row(s, 11.85, y, 6.4, 'check', a, b, circle=NAVY, d=0.62, label_size=17, desc_size=15, row_h=1.4)
notes(s, """
Walk through the five questions quickly. Everything today is illustrated with ONE running example from outpatient
care, so that the audience never has to hold two stories in their head. The card on the right is the contract with the
room: three things they will be able to do. Mention that the third one feeds directly into the use-case assessment
session that follows at 11:30.
""")

# 3 ---- Recap bridge ------------------------------------------------------------------------
s = content_slide('From Session 1: an agent in one sentence', 'Perceive, reason, act, and keep going until the goal is met')
c = rect(s, X0, 2.95, XW, 2.55, fill=LIGHT, radius=0.25)
text(s, X0 + 0.5, 3.1, XW - 1.0, 2.25,
     'An AI agent is a system that pursues a **goal** by reasoning about what to do next, using **tools** to act, '
     'and checking the result, inside **boundaries** set by people.', size=28, color=NAVY, anchor='m', line_spacing=1.1)
cols = [
    ('chat', 'Chatbot', 'Answers one question at a time. Stops when it has replied.', WHITE, INK, SLATE),
    ('cog', 'Automation', 'Follows a fixed script. Breaks when the situation is not in the script.', WHITE, INK, SLATE),
    ('robot', 'AI agent', 'Works towards a goal, step by step, choosing tools and adapting the plan as it goes.', BLUE, WHITE, WHITE),
]
cw = (XW - 2 * 0.45) / 3
for i, (ic, t, d, fill, tc, ic_col) in enumerate(cols):
    x = X0 + i * (cw + 0.45); y = 6.0; h = 3.9
    card(s, x, y, cw, h, fill=fill)
    variant = 'w' if fill == BLUE else 'b'
    icon_circle(s, x + 0.75, y + 0.75, 0.9, (NAVY if fill == WHITE else WHITE), ic, variant=('w' if fill == WHITE else 'b'), shadow=False)
    text(s, x + 1.45, y + 0.42, cw - 1.8, 0.6, t, size=24, color=tc, bold=True)
    text(s, x + 0.4, y + 1.55, cw - 0.8, h - 1.8, d, size=17, color=tc)
notes(s, """
One-slide recap so that everybody starts from the same definition. Three words carry the whole session: goal, tools,
boundaries. The comparison row is the litmus test: a chatbot answers, automation follows a script, an agent works
towards a goal and adapts. If you are short on time, skip this slide; the definition is repeated on slide 6.
""")

# 4 ---- Running example ---------------------------------------------------------------------
s = content_slide('Our running example: the No-Show Agent', 'One outpatient problem we will use on every slide today')
card(s, X0, 3.3, 8.3, 6.5)
icon_circle(s, X0 + 1.1, 3.3, 1.2, NAVY, 'hospital')
heading(s, X0 + 0.45, 4.15, 7.4, 'The situation', size=24)
body(s, X0 + 0.45, 4.75, 7.4, 4.9, [
    'Every day some patients do not turn up for their outpatient appointment.',
    {'t': 'Each missed slot is a wasted clinician hour, a longer waiting list, and a patient whose care is delayed.', 'space_before': 6},
    {'t': 'Reminder calls help, but staff cannot phone every patient every day.', 'space_before': 6},
], size=18)
card(s, X0 + 8.95, 3.3, 8.3, 6.5)
icon_circle(s, X0 + 8.95 + 1.1, 3.3, 1.2, BLUE, 'calendar')
heading(s, X0 + 9.4, 4.15, 7.4, 'What we want the agent to do', size=24)
body(s, X0 + 9.4, 4.75, 7.4, 4.9, [
    {'t': 'Spot who is likely to miss **tomorrow\'s** appointment', 'bullet': True},
    {'t': 'Reach out and confirm, or help them **reschedule**', 'bullet': True},
    {'t': 'Fill the freed slot from the **waiting list**', 'bullet': True},
    {'t': 'Hand over to staff when it is unsure or the case is sensitive', 'bullet': True},
    {'t': 'And never make a **clinical** decision', 'bullet': True, 'bullet_color': RED},
], size=18, space_after=8)
text(s, X0, 10.0, XW, 0.4, 'A fictional teaching scenario. The pattern applies to any clinic, any specialty, and most back-office processes.',
     size=13, color=SLATE, align='l')
notes(s, """
Set the scene in 60 seconds. Ask the room: who has seen a no-show problem in their clinic or department? Most hands
go up, which makes the rest of the session concrete. Stress the last bullet: the agent handles logistics, never
medicine. That single boundary will come back in every section. The scenario is fictional but realistic; invite people
to swap in their own process (referrals, discharge follow-up, prescription renewals) as we go.
""")

# 5 ---- Section: Anatomy --------------------------------------------------------------------
s = section_slide('Anatomy of an AI agent', 'Six parts every agent has, whatever it is built for', 'layers')
notes(s, "Section 1 of 5. About 14 minutes. One slide per building block, then the six blocks filled in for the No-Show Agent.")

# 6 ---- Anatomy overview --------------------------------------------------------------------
s = content_slide('Anatomy of an AI agent', 'Six building blocks, and the human is one of them')
cx, cy = 10.0, 6.75
blocks = [
    ('goal', '1  Goal', 'What it must achieve, and what it must never do'),
    ('planning', '2  Planning', 'How it breaks the goal into steps and adapts them'),
    ('knowledge', '3  Knowledge', 'What it knows: data, documents, memory'),
    ('tools', '4  Tools', 'What it can use: systems, models, rules, channels'),
    ('actions', '5  Actions', 'What it actually does in the real world'),
    ('oversight', '6  Human oversight', 'Where people set limits, approve and review'),
]
cw_, ch_ = 6.2, 2.05
pos = [(X0, 3.0), (X0, 5.55), (X0, 8.1), (X1 - cw_, 3.0), (X1 - cw_, 5.55), (X1 - cw_, 8.1)]
for (ic, lab, desc), (x, y) in zip(blocks, pos):
    left = x < 9
    # dotted connector to the centre
    connector(s, (x + cw_) if left else x, y + ch_ / 2, cx + (-1.55 if left else 1.55), cy, color=MID, width=1.5, dash=MSO_LINE.ROUND_DOT)
    card(s, x, y, cw_, ch_)
    icon_circle(s, x + 0.95, y + ch_ / 2, 1.1, BLUE if ic != 'oversight' else NAVY, ic, shadow=False)
    text(s, x + 1.75, y + 0.33, cw_ - 2.0, 0.5, lab, size=22, color=INK, bold=True)
    text(s, x + 1.75, y + 0.88, cw_ - 2.0, 1.0, desc, size=15, color=INK)
oval(s, cx, cy, 3.1, fill=NAVY, shadow=True)
oval(s, cx, cy, 2.75, fill=None, line=WHITE, line_w=1.5)
text(s, cx - 1.3, cy - 0.7, 2.6, 1.4, ['AI AGENT', {'t': 'a system, not a model', 'size': 13}], size=22, color=WHITE, bold=True, align='c', anchor='m')
notes(s, """
The map for the next six slides. Key message, borrowed from Session 1: an agent is a system, not a model. The language
model is only one part (it sits mostly in Planning). Point out that Human oversight is drawn as a building block, not
as something bolted on afterwards; that is deliberate and is the theme of the last two sections. Read the six blocks
aloud with the No-Show Agent in mind, then go one by one.
""")

# 7 ---- Goal --------------------------------------------------------------------------------
s = content_slide('1 · Goal', 'What the agent is trying to achieve, and what it must never do')
card(s, X0, 3.0, 5.6, 7.15)
heading(s, X0 + 0.4, 3.35, 4.9, 'A good goal is…', size=22, color=BLUE)
rows = [('goal', 'Specific', 'Which outcome, for whom, by when'),
        ('gauge', 'Measurable', 'How we will know the agent succeeded'),
        ('fence', 'Bounded', 'What it must never do, written into the goal itself')]
for i, (ic, a, b) in enumerate(rows):
    icon_row(s, X0 + 0.4, 4.2 + i * 1.85, 4.9, ic, a, b, circle=NAVY, d=0.66, label_size=18, desc_size=15, row_h=1.5)
# weak vs strong
gx = X0 + 6.1; gw = (X1 - gx - 0.45) / 2
card(s, gx, 3.0, gw, 5.5)
icon_circle(s, gx + 0.7, 3.65, 0.8, RED, 'x', shadow=False, scale=0.5)
text(s, gx + 1.25, 3.4, gw - 1.5, 0.5, 'Weak goal', size=20, color=INK, bold=True)
text(s, gx + 0.4, 4.45, gw - 0.8, 3.8, ['“Help reduce no-shows.”',
     {'t': 'Nothing says which patients, which clinics, what counts as success, or what the agent may not do. '
           'A goal like this leaves the model to guess, and it will.', 'size': 15, 'space_before': 10}],
     size=22, color=NAVY)
gx2 = gx + gw + 0.45
card(s, gx2, 3.0, gw, 5.5)
icon_circle(s, gx2 + 0.7, 3.65, 0.8, GREEN, 'check', shadow=False, scale=0.5)
text(s, gx2 + 1.25, 3.4, gw - 1.5, 0.5, 'Strong goal', size=20, color=INK, bold=True)
text(s, gx2 + 0.4, 4.45, gw - 0.8, 3.9, [
    '“Every working day, contact outpatients at high risk of missing tomorrow\'s appointment, confirm or reschedule, '
    'and refill freed slots from the waiting list.',
    {'t': 'Never cancel without the patient\'s consent. Never give medical advice.”', 'space_before': 6, 'bold': True}],
    size=17, color=NAVY)
callout(s, gx, 8.85, X1 - gx, 1.3,
        'The goal is written by **people**, not by the model. Think of it as the agent\'s job description.', icon_name='users')
notes(s, """
Ask: what would you tell a new colleague on their first day? That is the goal. Read the weak goal, then the strong one.
The strong goal contains three things: the outcome, the scope, and the must-never rules. Non-technical people are the
best authors of goals because they know the process. Practical tip: if the team cannot write the goal in three
sentences, the use case is not ready for an agent yet. This links to the use-case assessment session after the break.
""")

# 8 ---- Planning ----------------------------------------------------------------------------
s = content_slide('2 · Planning', 'How the agent turns a goal into steps, and changes the plan when reality changes')
card(s, X0, 3.0, 5.7, 7.15)
icon_circle(s, X0 + 0.85, 3.85, 1.0, BLUE, 'brain', shadow=False)
text(s, X0 + 1.6, 3.5, 3.9, 0.7, 'The reasoning engine plans', size=20, color=INK, bold=True)
body(s, X0 + 0.4, 4.75, 4.9, 5.2, [
    'Given the goal, the current situation and the tools it may use, the language model proposes the **next step**.',
    {'t': 'It does not follow a fixed script. When something unexpected happens it **re-plans**.', 'space_before': 8},
    {'t': 'That flexibility is the difference from classic automation, and the reason we need boundaries (section 5).',
     'space_before': 8, 'color': BLUE},
], size=16)
px = X0 + 6.2; pw = X1 - px
heading(s, px, 2.95, pw, 'The No-Show Agent\'s plan for tomorrow\'s clinic', size=20, color=BLUE)
steps = [('Get tomorrow\'s appointment list', 'from the scheduling system'),
         ('Score each patient\'s no-show risk', 'using the prediction model'),
         ('Contact high-risk patients', 'with an approved message, offering to confirm or reschedule'),
         ('If they cannot come, offer alternative slots', 'within the clinic\'s rules'),
         ('Fill the freed slot from the waiting list', 'and confirm to everyone involved')]
for i, (a, b) in enumerate(steps):
    y = 3.6 + i * 0.98
    numbered_row(s, px, y, pw, i + 1, a, b, d=0.6, label_size=17, desc_size=14, row_h=0.9)
callout(s, px, 8.6, pw, 1.55, [
    {'t': 'When reality changes, the plan changes', 'bold': True, 'size': 17},
    {'t': 'No reply after two messages → try a phone call → still nothing → hand the case to the clinic coordinator.', 'size': 15}],
    icon_name='loop')
notes(s, """
Planning is where the language model earns its place: it turns a goal into a sequence of steps and revises it when the
world does not cooperate. Walk the five steps; they read like a checklist a good coordinator would write. Then the
callout: the plan bends. That is powerful and slightly scary, which is exactly why sections 4 and 5 exist. Note for the
curious: in practice the plan can be partly fixed by the designer (a workflow) and partly left to the model; the
hands-on session this afternoon shows both.
""")

# 9 ---- Knowledge ---------------------------------------------------------------------------
s = content_slide('3 · Knowledge', 'An agent is only as smart as the resources available to it')
kcols = [
    ('brain', 'General knowledge', 'From the model\'s training: language, common sense, how a conversation works.',
     'Cannot know EHS policies or today\'s schedule.', NAVY),
    ('files', 'Enterprise documents', 'Retrieved when needed: scheduling policy, clinic rules, approved message templates.',
     'Grounds answers in your own knowledge (RAG).', BLUE),
    ('database', 'Live data', 'Read from systems at run time: tomorrow\'s list, contact preferences, waiting list.',
     'Always current, never copied into the model.', SKY),
    ('memory', 'Memory', 'What happened so far in this task: “this patient already declined twice this month.”',
     'Keeps the conversation coherent.', BLUE),
]
kw = (XW - 3 * 0.42) / 4
for i, (ic, t, d, foot, col) in enumerate(kcols):
    x = X0 + i * (kw + 0.42); y = 3.45; h = 5.35
    card(s, x, y, kw, h)
    icon_circle(s, x + kw / 2, y, 1.15, col, ic)
    text(s, x + 0.3, y + 0.85, kw - 0.6, 0.6, t, size=20, color=INK, bold=True, align='c')
    text(s, x + 0.35, y + 1.6, kw - 0.7, 2.3, d, size=15, color=INK)
    text(s, x + 0.35, y + h - 1.2, kw - 0.7, 0.9, foot, size=14, color=BLUE, italic=False) if False else \
        text(s, x + 0.35, y + h - 1.25, kw - 0.7, 1.0, [{'t': foot, 'color': BLUE, 'size': 14}])
callout(s, X0, 9.15, XW, 1.05, 'Rule of thumb: if the answer is not in one of these four places, the agent will **guess**. '
        'Good agent design removes the need to guess.', icon_name='lightbulb')
notes(s, """
Four sources of knowledge, in order of trust: what the model learned in training (general, sometimes outdated),
enterprise documents retrieved on demand (that is RAG from the pre-reading deck), live data from systems, and memory of
the current task. The failure mode to name explicitly: when none of these contains the answer, the model produces a
plausible guess. That is the hallucination problem, and the cure is design: give the agent the documents and data
it needs, and tell it to say 'I do not know' otherwise.
""")

# 10 ---- Tools ------------------------------------------------------------------------------
s = content_slide('4 · Tools', 'What the agent can pick up to get the job done')
tcols = [
    ('database', 'Data queries', 'Read appointments, contact details, waiting list'),
    ('model', 'AI / ML models', 'No-show risk score, readmission risk, forecasts'),
    ('rules', 'Business rules', 'Which appointment types may be moved, and how far'),
    ('api', 'Applications & channels', 'WhatsApp / SMS, e-mail, booking system, ticket queue'),
    ('agents', 'Other agents', 'A waiting-list agent, a translation agent'),
]
tw = (XW - 4 * 0.35) / 5
for i, (ic, t, d) in enumerate(tcols):
    x = X0 + i * (tw + 0.35); y = 3.45; h = 4.25
    card(s, x, y, tw, h)
    icon_circle(s, x + tw / 2, y, 1.05, NAVY if i % 2 == 0 else BLUE, ic)
    text(s, x + 0.2, y + 0.8, tw - 0.4, 0.95, t, size=18, color=INK, bold=True, align='c')
    text(s, x + 0.3, y + 1.8, tw - 0.6, 2.3, d, size=14, color=INK, align='c')
card(s, X0, 8.15, XW, 2.05, fill=PALE, shadow=False)
icon(s, 'key', 'n', X0 + 0.45, 8.65, 1.0)
text(s, X0 + 1.85, 8.35, XW - 2.2, 1.7, [
    {'t': 'The model never runs a tool itself. It **asks**; the agent runtime **executes**, with the permissions we granted.', 'size': 18},
    {'t': 'Tools are described to the model in plain language (“check free slots for a clinic and date”). '
          'Standards such as MCP make tools plug-and-play; you will connect one yourself in Session 5 this afternoon.',
     'size': 14, 'color': SLATE, 'space_before': 6}],
    color=NAVY, anchor='m')
notes(s, """
A tool is anything the agent can call: a query, a model, a rule set, an application, even another agent. Two points
matter for a non-technical audience. First, the model does not execute anything; it requests, and a runtime that we
control performs the action under our permissions. That is where security lives (Session 4). Second, tools are
described in plain language, so business people can review the tool list and understand what the agent can reach.
MCP was covered in the pre-read deck as the USB-C of agents; the hands-on this afternoon makes it tangible.
""")

# 11 ---- Actions ----------------------------------------------------------------------------
s = content_slide('5 · Actions', 'Not all actions are equal: reading is cheap, writing is a commitment')
arrow(s, X0, 3.0, XW, 0.62, fill=LIGHT)
text(s, X0 + 0.4, 3.0, XW - 1.2, 0.62, 'Increasing impact  →  stronger controls', size=16, color=NAVY, bold=True, anchor='m')
acols = [
    ('READ', 'Look up tomorrow\'s appointments', 'No side effects. The world is unchanged.', LIGHT, NAVY, 'Free to run'),
    ('SUGGEST', 'Draft a reminder message for a person to send', 'Nothing happens until a human acts.', SKY, WHITE, 'Human sends'),
    ('ACT · reversible', 'Send a reminder, book a slot', 'Can be undone. A mistake costs a little time.', BLUE, WHITE, 'Log it, monitor it'),
    ('ACT · irreversible', 'Cancel an appointment, send a clinical instruction', 'Cannot be undone. A mistake costs trust, or worse.', NAVY, WHITE, 'Approval or human only'),
]
aw = (XW - 3 * 0.4) / 4
for i, (t, ex, d, col, tc, ctrl) in enumerate(acols):
    x = X0 + i * (aw + 0.4); y = 4.0; h = 5.0
    card(s, x, y, aw, h)
    p = rect(s, x + 0.3, y + 0.35, aw - 0.6, 0.75, fill=col, radius=0.375)
    shape_text(p, t, size=16, color=tc, bold=True)
    text(s, x + 0.35, y + 1.35, aw - 0.7, 1.6, ex, size=17, color=INK, bold=True)
    text(s, x + 0.35, y + 2.85, aw - 0.7, 1.3, d, size=15, color=INK)
    text(s, x + 0.35, y + h - 0.85, aw - 0.7, 0.6, [{'t': ctrl, 'color': BLUE, 'size': 14, 'bold': True}])
callout(s, X0, 9.3, XW, 0.9, 'We will use this ladder in section 4 to decide how much autonomy each action gets.', icon_name='sliders', size=16)
notes(s, """
This is the most useful mental model of the day for non-technical decision makers. Sort every action the agent could
take into these four buckets. Reads are free. Suggestions are safe because a person acts. Reversible actions can be
automated with monitoring. Irreversible actions need approval or stay with humans. Ask the room to classify: sending a
WhatsApp reminder (reversible), moving an oncology appointment (irreversible in effect), telling a patient a result
(clinical, human only).
""")

# 12 ---- Human oversight --------------------------------------------------------------------
s = content_slide('6 · Human oversight', 'People stay in charge at three moments')
ocols = [
    ('clipboard', 'Before', 'Design time', ['Write the goal and the must-never rules', 'Choose the tools and the permissions',
                                            'Set the limits and the escalation routes', 'Test in shadow mode']),
    ('oversight', 'During', 'Run time', ['Approve sensitive actions', 'Answer escalations', 'Take over a case at any time',
                                         'Pause the agent with one switch']),
    ('log', 'After', 'Review', ['Read the logs: what it did and why', 'Audit outcomes and exceptions',
                                'Tune goal, rules and tools', 'Decide whether to grant more autonomy']),
]
ow = (XW - 2 * 0.45) / 3
for i, (ic, t, sub, bl) in enumerate(ocols):
    x = X0 + i * (ow + 0.45); y = 3.45; h = 5.55
    card(s, x, y, ow, h)
    icon_circle(s, x + ow / 2, y, 1.15, [NAVY, BLUE, SKY][i], ic)
    text(s, x + 0.3, y + 0.8, ow - 0.6, 0.55, t, size=24, color=INK, bold=True, align='c')
    text(s, x + 0.3, y + 1.35, ow - 0.6, 0.45, sub, size=15, color=SLATE, align='c')
    text(s, x + 0.5, y + 2.0, ow - 0.9, 3.4, [{'t': b, 'bullet': True} for b in bl], size=16, color=INK, space_after=7)
callout(s, X0, 9.3, XW, 0.95, 'Oversight is not a brake on the agent. It is what earns the agent more autonomy later.', icon_name='trophy', size=17)
notes(s, """
Oversight is a building block, not an afterthought. Three moments: before (design), during (operation), after
(review). Most of the room will be 'before' and 'after' people: they write goals and rules, and they read the reports.
The 'during' role is usually a coordinator with an approval queue. End with the callout: the better the oversight,
the more autonomy you can safely give, which is the trust ladder we will see in section 4.
""")

# 13 ---- Canvas -----------------------------------------------------------------------------
s = content_slide('The No-Show Agent on one page', 'The same six blocks, filled in. You will reuse this canvas in the use-case exercise.')
canvas = [
    ('goal', '1  Goal', 'Reduce missed outpatient appointments by confirming or rescheduling high-risk patients the day before, '
                        'and refilling freed slots. Never cancel without consent. Never give medical advice.'),
    ('planning', '2  Planning', 'Daily: list → risk score → contact → reschedule if needed → refill slot. '
                                'Re-plan on no reply; escalate after two failed attempts.'),
    ('knowledge', '3  Knowledge', 'Scheduling policy and message templates (documents); appointment list, contact preferences, '
                                  'waiting list (live data); attempt history (memory).'),
    ('tools', '4  Tools', 'Scheduling system (read / write bookings), no-show risk model, WhatsApp & SMS gateway, '
                          'clinic rule set, coordinator task queue.'),
    ('actions', '5  Actions', 'Send approved reminders; propose and book alternative slots with patient consent; '
                              'offer freed slots to the waiting list; create hand-over tasks.'),
    ('oversight', '6  Human oversight', 'Coordinator approves moves of specialist appointments and handles escalations; '
                                        'weekly review of logs, exceptions and no-show rate.'),
]
gw = (XW - 2 * 0.4) / 3; gh = 3.55
for i, (ic, t, d) in enumerate(canvas):
    r, c = divmod(i, 3)
    x = X0 + c * (gw + 0.4); y = 2.95 + r * (gh + 0.4)
    card(s, x, y, gw, gh)
    icon_circle(s, x + 0.65, y + 0.62, 0.72, NAVY if ic == 'oversight' else BLUE, ic, shadow=False, scale=0.55)
    text(s, x + 1.2, y + 0.35, gw - 1.45, 0.55, t, size=20, color=INK, bold=True)
    text(s, x + 0.4, y + 1.2, gw - 0.8, gh - 1.4, d, size=15, color=INK, line_spacing=1.05)
notes(s, """
Everything from the last six slides on one page. This is the deliverable of the section: an 'agent canvas' that a
business team can fill in without any technical help. Tell the audience they will fill one in for their own use case
in the assessment session that follows, and again on Day 2. Pause for questions here before moving to the loop.
""")

# 14 ---- Section: loop ----------------------------------------------------------------------
s = section_slide('How an agent decides and executes', 'The loop that makes it an agent', 'loop', title_size=54)
notes(s, "Section 2 of 5. About 9 minutes: the loop, one worked case, what happens inside a reasoning step, and how the loop stops.")

# 15 ---- The loop ---------------------------------------------------------------------------
s = content_slide('The agent loop: think, act, observe, repeat', 'A chatbot answers once. An agent keeps going until the goal is met or a boundary is hit.')
lcx, lcy, R = 5.95, 6.75, 2.35
ring = slide_ring = s.shapes.add_shape(MSO_SHAPE.DONUT, Inches(lcx - R - 0.42), Inches(lcy - R - 0.42), Inches(2 * R + 0.84), Inches(2 * R + 0.84))
ring.adjustments[0] = 0.15; _fill(ring, LIGHT); _no_line(ring); _no_shadow(ring)
nodes = [('PERCEIVE', 'What is the situation now?', 270, NAVY), ('REASON & PLAN', 'What is the best next step?', 0, BLUE),
         ('ACT', 'Use a tool, send a message', 90, NAVY), ('OBSERVE', 'What happened? Did it work?', 180, BLUE)]
for lab, d, ang, col in nodes:
    a = math.radians(ang); nx = lcx + R * math.cos(a); ny = lcy + R * math.sin(a)
    o = oval(s, nx, ny, 1.6, fill=col, shadow=True)
    shape_text(o, lab, size=13, color=WHITE, bold=True, margins=(0.05, 0.02, 0.05, 0.02))
for phi in (315, 45, 135, 225):
    a = math.radians(phi); px_ = lcx + R * math.cos(a); py_ = lcy + R * math.sin(a)
    chevron(s, px_ - 0.3, py_ - 0.26, 0.6, 0.52, fill=BLUE, rotation=phi + 90)
text(s, lcx - 1.35, lcy - 0.85, 2.7, 1.7, ['Repeat until the goal is met,', {'t': 'or a boundary is hit', 'color': BLUE}], size=14, color=NAVY, bold=True, align='c', anchor='m')
text(s, lcx - 1.6, lcy + R + 0.9, 3.2, 0.5, '', size=10)
rx = 10.0; rw = X1 - rx
card(s, rx, 3.0, rw, 5.35)
heading(s, rx + 0.45, 3.3, rw - 0.9, 'Each turn of the loop', size=20, color=BLUE)
lrows = [('eye', 'Perceive', 'Read the situation: new appointment list, a patient reply, a tool result.'),
         ('brain', 'Reason & plan', 'The model weighs goal, rules, knowledge and tools, and proposes the next step.'),
         ('play', 'Act', 'The runtime executes the step: a query, a message, a booking.'),
         ('search', 'Observe', 'The result comes back and becomes the next thing to perceive.')]
for i, (ic, a, b) in enumerate(lrows):
    icon_row(s, rx + 0.45, 3.95 + i * 1.08, rw - 0.9, ic, a, b, circle=NAVY if i % 2 == 0 else BLUE, d=0.58, label_size=16, desc_size=14, row_h=1.0)
callout(s, rx, 8.65, rw, 1.55, [{'t': 'Guardrail inside the loop', 'bold': True, 'size': 16},
        {'t': 'Before any action runs, it is checked against the boundaries (section 5). The model proposes; the rules dispose.', 'size': 14}],
        icon_name='shield')
notes(s, """
This loop is the definition of 'agentic'. Four beats: perceive, reason and plan, act, observe, then round again. A
chatbot runs the loop once; an agent runs it until the goal is met or a boundary stops it. The guardrail callout is the
bridge to the security session: every proposed action is checked against the rules before it executes. If people ask
'where is the LLM?', point to Reason & plan: that is where it lives, and only there.
""")

# 16 ---- Walkthrough ------------------------------------------------------------------------
s = content_slide('The loop in practice', 'One patient, one afternoon: watch the No-Show Agent work through a single case')
rect(s, X0, 2.85, XW, 7.45, fill=PANEL, radius=0.35, shadow=True)
wsteps = [('1 · Perceive', 'Tomorrow\'s cardiology list arrives: 42 appointments. The risk model flags Mr A: two no-shows this year, a long commute.'),
          ('2 · Reason & plan', 'Goal says: contact high-risk patients. Policy allows an approved reminder template. Plan: send it, with a reschedule option.'),
          ('3 · Act', 'WhatsApp message sent through the approved template, in the patient\'s preferred language.'),
          ('4 · Observe', 'Reply: “I can\'t make it, I have no transport that day.”'),
          ('5 · Reason & plan', 'Routine follow-up: policy allows moving it within 14 days. Two free slots next week. Propose both; do not cancel.'),
          ('6 · Act → done', 'Patient picks Tuesday. Slot rebooked, confirmation sent, freed slot offered to the first waiting-list patient. Everything logged.')]
cw2 = 5.05; gap = 0.67; ch2 = 2.75
for i, (t, d) in enumerate(wsteps):
    r, c = divmod(i, 3)
    x = X0 + 0.5 + c * (cw2 + gap); y = 3.65 + r * (ch2 + 0.95)
    card(s, x, y, cw2, ch2, shadow=False)
    p = rect(s, x + 0.35, y - 0.3, cw2 - 0.7, 0.6, fill=NAVY, radius=0.3, shadow=True)
    shape_text(p, t, size=15, color=WHITE, bold=True, margins=(0.1, 0.02, 0.1, 0.02))
    text(s, x + 0.35, y + 0.5, cw2 - 0.7, ch2 - 0.65, d, size=14, color=INK, anchor='m', align='c')
    if c < 2:
        arrow(s, x + cw2 + 0.12, y + ch2 / 2 - 0.22, gap - 0.24, 0.44, fill=BLUE)
notes(s, """
Tell it as a story. Six beats, two full turns of the loop plus a close. Point out three design decisions hiding in the
story: the message came from an approved template (boundary), the agent proposed slots rather than cancelling (hard
rule), and everything was logged (oversight after the fact). Ask: where could this have gone wrong? Typical answers:
patient asks a medical question, no free slot within policy, patient does not reply. Each is an escalation point,
which we cover in section 5.
""")

# 17 ---- Inside one reasoning step ---------------------------------------------------------
s = content_slide('Inside one reasoning step', 'What the model is given, and what it gives back')
lw = 6.0
card(s, X0, 3.0, lw, 5.75)
heading(s, X0 + 0.4, 3.3, lw - 0.8, 'What goes in', size=20, color=BLUE)
inrows = [('goal', 'Goal & instructions', 'The job description (the system prompt)'),
          ('files', 'The situation', 'This case, the conversation so far, tool results'),
          ('tools', 'Available tools', 'Each with a plain-language description'),
          ('fence', 'Rules & limits', 'What it may and may not do')]
for i, (ic, a, b) in enumerate(inrows):
    icon_row(s, X0 + 0.4, 3.95 + i * 1.15, lw - 0.8, ic, a, b, circle=NAVY, d=0.58, label_size=16, desc_size=14, row_h=1.05)
mcx = 10.0; mcy = 5.9
connector(s, X0 + lw + 0.1, mcy, mcx - 1.5, mcy, color=BLUE, width=2.25, tail=True)
oval(s, mcx, mcy, 2.8, fill=BLUE, shadow=True)
oval(s, mcx, mcy, 2.45, fill=None, line=WHITE, line_w=1.5)
icon(s, 'brain', 'w', mcx - 0.45, mcy - 0.95, 0.9)
text(s, mcx - 1.2, mcy - 0.05, 2.4, 1.1, ['Reasoning engine', {'t': '(the language model)', 'size': 12}], size=15, color=WHITE, bold=True, align='c', anchor='t')
rx = X1 - lw
connector(s, mcx + 1.5, mcy, rx - 0.1, mcy, color=BLUE, width=2.25, tail=True)
card(s, rx, 3.0, lw, 5.75)
heading(s, rx + 0.4, 3.3, lw - 0.8, 'What comes out: one of three', size=20, color=BLUE)
outrows = [('tools', 'A tool request', '“Check free cardiology slots next week”'),
           ('message', 'A message', 'To the patient, or to a colleague'),
           ('hand', 'A hand-over', '“This needs a person”, with the reason')]
for i, (ic, a, b) in enumerate(outrows):
    icon_row(s, rx + 0.4, 3.95 + i * 1.45, lw - 0.8, ic, a, b, circle=BLUE, d=0.58, label_size=16, desc_size=14, row_h=1.3)
callout(s, X0, 9.15, XW, 1.05, 'Analogy: a capable new colleague with a job description, the policy binder, a phone list, '
        'and a supervisor they can call. The runtime, not the colleague, makes the calls.', icon_name='users', size=16)
notes(s, """
Demystify the 'thinking'. Every turn, the model receives a package: the goal and instructions, the situation so far,
the list of tools with descriptions, and the rules. It returns exactly one of three things: a tool request, a message,
or a hand-over. It cannot do anything else. That is why the tool list and the rules are the real levers of control.
The analogy of the new colleague works well with non-technical audiences; use it.
""")

# 18 ---- When the loop stops ----------------------------------------------------------------
s = content_slide('When does the loop stop?', 'Agents do not run forever. We design the exits.')
ecols = [('check', 'Goal achieved', 'Appointment confirmed or rebooked; freed slot refilled.', 'Log the outcome, close the case.', GREEN),
         ('hand', 'Hand-over', 'An escalation trigger fired: low confidence, out of scope, sensitive case.', 'Task to a named person, with full context.', BLUE),
         ('clock', 'Limit reached', 'Two contact attempts made, or the 20:00 messaging cut-off passed.', 'Stop, report, retry tomorrow if allowed.', NAVY),
         ('stop', 'Blocked', 'A tool is down, or two rules conflict and the agent cannot proceed.', 'Stop safely, alert the owner.', SLATE)]
ew = (XW - 3 * 0.4) / 4
for i, (ic, t, d, nxt, col) in enumerate(ecols):
    x = X0 + i * (ew + 0.4); y = 3.45; h = 5.35
    card(s, x, y, ew, h)
    icon_circle(s, x + ew / 2, y, 1.1, col, ic)
    text(s, x + 0.3, y + 0.8, ew - 0.6, 0.6, t, size=20, color=INK, bold=True, align='c')
    text(s, x + 0.35, y + 1.55, ew - 0.7, 2.1, d, size=15, color=INK)
    rect(s, x + 0.35, y + h - 1.75, ew - 0.7, 1.35, fill=PALE, radius=0.15)
    text(s, x + 0.5, y + h - 1.7, ew - 1.0, 1.25, [{'t': 'Then: ', 'bold': True, 'color': BLUE, 'size': 14}, {'t': nxt, 'size': 14, 'color': NAVY}], anchor='m')
callout(s, X0, 9.15, XW, 1.05, 'Every exit is logged and visible. Silence is never an acceptable outcome for an agent.', icon_name='log', size=17)
notes(s, """
A common worry: 'does it just keep going?'. No. Four exits, all designed by us: success, hand-over, a limit, or a
block. Limits are boring but essential: attempts, time windows, daily volumes, budgets. Blocks are the safe failure:
stop and alert rather than improvise. The rule at the bottom is a good one to write on the wall: silence is never an
acceptable outcome.
""")

# 19 ---- Section: enterprise ----------------------------------------------------------------
s = section_slide('Connecting to the enterprise', 'Systems, data and tools, behind a governed front door', 'network')
notes(s, "Section 3 of 5. About 5 minutes: the integration picture and four habits that keep it safe.")

# 20 ---- Architecture -----------------------------------------------------------------------
s = content_slide('Reaching enterprise systems, data and tools', 'Everything passes through a governed tool layer, never a direct line into the database', title_size=52)
# trigger
icon_circle(s, 2.2, 6.3, 1.5, SLATE, 'clock', shadow=True)
text(s, 1.1, 7.2, 2.2, 0.9, ['07:00 daily run', {'t': 'or a patient reply', 'size': 12}], size=14, color=INK, bold=True, align='c')
connector(s, 3.05, 6.3, 3.85, 6.3, color=SLATE, width=2, tail=True)
# agent
ax, ay, aw_, ah = 4.0, 3.5, 4.9, 5.9
card(s, ax, ay, aw_, ah)
icon_circle(s, ax + aw_ / 2, ay, 1.15, NAVY, 'robot')
text(s, ax + 0.3, ay + 0.8, aw_ - 0.6, 0.6, 'AI agent', size=22, color=INK, bold=True, align='c')
arows = [('goal', 'Goal & rules'), ('planning', 'Plan'), ('memory', 'Memory'), ('brain', 'Reasoning engine')]
for i, (ic, t) in enumerate(arows):
    yy = ay + 1.65 + i * 0.98
    icon_circle(s, ax + 0.75, yy + 0.3, 0.6, BLUE, ic, shadow=False, scale=0.55)
    text(s, ax + 1.25, yy + 0.05, aw_ - 1.5, 0.5, t, size=16, color=INK, bold=True)
# governed layer
gx, gy, gw2, gh2 = 9.6, 2.95, 2.4, 7.3
g = rect(s, gx, gy, gw2, gh2, fill=NAVY, radius=0.3, shadow=True)
text(s, gx - 2.3, gy + gh2 / 2 - 0.9, gh2 - 0.4, 1.8, ['GOVERNED TOOL LAYER', {'t': 'identity  ·  permissions  ·  audit log  ·  MCP', 'size': 13, 'bold': False}],
     size=18, color=WHITE, bold=True, align='c', anchor='m', rotation=270)
# reposition rotated textbox: python-pptx rotates around centre, so place its centre on the bar centre
tb = s.shapes[-1]; tb.left = Inches(gx + gw2 / 2 - (gh2 - 0.4) / 2); tb.top = Inches(gy + gh2 / 2 - 0.9)
connector(s, ax + aw_ + 0.1, 6.3, gx - 0.1, 6.3, color=BLUE, width=2.25, tail=True, head=True)
# right stack
rcols = [('database', 'Enterprise data', 'Scheduling system, patient contact details, waiting list'),
         ('files', 'Documents & policies', 'Scheduling policy, message templates (retrieval / RAG)'),
         ('model', 'Models & rules', 'No-show risk model, decision rules (SAS Viya)'),
         ('message', 'Applications & channels', 'WhatsApp / SMS gateway, e-mail, coordinator task queue'),
         ('agents', 'Other agents', 'Waiting-list agent, translation agent')]
sx = 12.9; sw = X1 - sx; sh = 1.22; sgap = 0.3
for i, (ic, t, d) in enumerate(rcols):
    yy = gy + i * (sh + sgap)
    card(s, sx, yy, sw, sh)
    icon_circle(s, sx + 0.7, yy + sh / 2, 0.72, BLUE if i % 2 == 0 else SKY, ic, shadow=False, scale=0.55)
    text(s, sx + 1.25, yy + 0.12, sw - 1.45, 0.45, t, size=16, color=INK, bold=True)
    text(s, sx + 1.25, yy + 0.55, sw - 1.45, 0.6, d, size=12, color=SLATE)
    connector(s, gx + gw2 + 0.05, yy + sh / 2, sx - 0.08, yy + sh / 2, color=BLUE, width=1.75, tail=True)
notes(s, """
The picture to remember: the agent never touches a database directly. It talks to a governed tool layer that knows who
the agent is, what it may call, and writes every call to an audit log. Behind that layer sit the things the agent
needs: data, documents, models and rules, applications and channels, and sometimes other agents. In SAS Viya the
risk model and the decision rules sit on the right; the MCP server is the front door. Keep this slide short; Session 4
covers identity and access in depth.
""")

# 21 ---- Principles -------------------------------------------------------------------------
s = content_slide('What good integration looks like', 'Four habits that keep agents safe to connect')
pcols = [('key', 'Its own identity, least privilege', 'The agent logs in as itself with only the access it needs. The No-Show Agent can read appointments and write bookings; it cannot open clinical notes.'),
         ('shield', 'Validated tools, not open access', 'Each tool is a tested, versioned capability with a clear description. No free-form access to systems.'),
         ('database', 'Data stays where it is', 'The agent asks the system a question at run time. It never receives a copy of the database.'),
         ('log', 'Everything is logged', 'Every tool call and every message: who, what, why, when. Logs feed the review and the audit.')]
pw2 = (XW - 3 * 0.4) / 4
for i, (ic, t, d) in enumerate(pcols):
    x = X0 + i * (pw2 + 0.4); y = 3.45; h = 5.35
    card(s, x, y, pw2, h)
    icon_circle(s, x + pw2 / 2, y, 1.1, NAVY if i % 2 == 0 else BLUE, ic)
    text(s, x + 0.3, y + 0.8, pw2 - 0.6, 1.1, t, size=18, color=INK, bold=True, align='c')
    text(s, x + 0.35, y + 2.0, pw2 - 0.7, 3.2, d, size=14, color=INK)
callout(s, X0, 9.15, XW, 1.05, 'Session 4 this afternoon goes deeper: agent identity, access management and the risks of autonomous agents.', icon_name='lock', size=16)
notes(s, """
Four habits, no jargon. Own identity and least privilege: the agent is a user like any other, with a narrow badge.
Validated tools: you review a short list of named capabilities, not an API. Data stays put. Everything is logged.
These four make the difference between an agent you can defend in an audit and one you cannot. Hand over to Session 4
for the details; here we only need the principles.
""")

# 22 ---- Section: HITL ----------------------------------------------------------------------
s = section_slide('Human-in-the-loop vs autonomous', 'Same agent, different levels of trust', 'oversight', title_size=54)
notes(s, "Section 4 of 5. About 7 minutes: the autonomy spectrum, how to choose a level per action, and the trust ladder.")

# 23 ---- Spectrum ---------------------------------------------------------------------------
s = content_slide('The autonomy spectrum', 'Four ways to run the same agent, from assistant to autonomous')
arrow(s, X0, 2.95, XW, 0.6, fill=LIGHT)
text(s, X0 + 0.4, 2.95, XW - 1.2, 0.6, 'More autonomy  →  more speed, more trust required', size=16, color=NAVY, bold=True, anchor='m')
levels = [('1 · Assist', 'Human in the loop', LIGHT, NAVY,
           'Recommends. Prepares a list of at-risk patients and a draft message for each.', 'Reviews, edits, sends. Owns every action.',
           'New use cases, sensitive processes, first weeks of a pilot.'),
          ('2 · Approve', 'Human in the loop', SKY, WHITE,
           'Proposes each action and waits. “Send this reminder to Mr A?”', 'Approves or rejects each action from a queue.',
           'Reversible actions with real impact; building trust.'),
          ('3 · Supervise', 'Human on the loop', BLUE, WHITE,
           'Acts on its own inside the rules. Flags exceptions.', 'Handles exceptions, reads a daily summary, can pause.',
           'Proven, reversible actions at volume: reminders, rebooking.'),
          ('4 · Autonomous', 'Human over the loop', NAVY, WHITE,
           'Acts end-to-end within hard boundaries.', 'Audits samples and outcomes; sets and reviews the boundaries.',
           'Low-risk, high-volume, well-measured actions only.')]
lw2 = (XW - 3 * 0.35) / 4
for i, (t, sub, col, tc, ag, hu, fit) in enumerate(levels):
    x = X0 + i * (lw2 + 0.35); y = 3.9; h = 6.3
    card(s, x, y, lw2, h)
    p = rect(s, x + 0.3, y + 0.3, lw2 - 0.6, 0.72, fill=col, radius=0.36)
    shape_text(p, t, size=17, color=tc, bold=True)
    text(s, x + 0.3, y + 1.08, lw2 - 0.6, 0.4, sub, size=13, color=SLATE, align='c')
    yy = y + 1.6
    for lab, txt_, ic in (('Agent', ag, 'robot'), ('Human', hu, 'user'), ('Fits when', fit, 'check')):
        icon(s, ic, 'b', x + 0.35, yy + 0.02, 0.34)
        text(s, x + 0.8, yy - 0.05, lw2 - 1.1, 0.4, lab, size=14, color=BLUE, bold=True)
        text(s, x + 0.35, yy + 0.38, lw2 - 0.7, 1.1, txt_, size=13, color=INK)
        yy += 1.5
notes(s, """
Four operating modes for the very same agent. Assist: it recommends, people act. Approve: it proposes, people click.
Supervise: it acts, people handle exceptions and read summaries. Autonomous: it acts, people audit. The vocabulary
'in the loop / on the loop / over the loop' is common in the literature; use whichever the room prefers. Key point:
this is a dial, not a switch, and you set it per action, which is the next slide.
""")

# 24 ---- Choosing the level ----------------------------------------------------------------
s = content_slide('Choosing the level for each action', 'Two questions: how bad is a wrong action, and can it be undone?')
mx, my, mw, mh = 2.6, 3.15, 9.3, 6.3
qw = (mw - 0.3) / 2; qh = (mh - 0.3) / 2
quads = [(0, 0, 'High impact · easy to undo', '2 · Approve', 'Move a specialist appointment; contact a patient flagged as vulnerable', SKY, WHITE),
         (1, 0, 'High impact · hard to undo', 'Human decides', 'Cancel an oncology appointment; anything clinical; anything a patient could read as advice', NAVY, WHITE),
         (0, 1, 'Low impact · easy to undo', '4 · Autonomous', 'Send an approved reminder; rebook a routine follow-up with consent', LIGHT, NAVY),
         (1, 1, 'Low impact · hard to undo', '3 · Supervise', 'Offer a freed slot to the waiting list; send a confirmation', BLUE, WHITE)]
for c, r, t, lvl, ex, col, tc in quads:
    x = mx + c * (qw + 0.3); y = my + r * (qh + 0.3)
    card(s, x, y, qw, qh, fill=col, shadow=True)
    text(s, x + 0.3, y + 0.25, qw - 0.6, 0.45, t, size=13, color=tc, bold=False)
    text(s, x + 0.3, y + 0.7, qw - 0.6, 0.6, lvl, size=22, color=tc, bold=True)
    text(s, x + 0.3, y + 1.4, qw - 0.6, qh - 1.6, ex, size=14, color=tc)
text(s, X0 - 0.2, my, 1.1, mh, 'Impact of a wrong action  →  higher', size=14, color=SLATE, bold=True, align='c', anchor='m', rotation=270)
tbx = s.shapes[-1]; tbx.left = Inches(mx - 0.55 - mh / 2 + 0.55 - 0.55); tbx.top = Inches(my + mh / 2 - 0.55)
tbx.left = Inches(mx - 0.75 - (mh - 1.1) / 2); tbx.width = Inches(mh); tbx.height = Inches(1.1)
text(s, mx, my + mh + 0.1, mw, 0.5, 'Harder to undo  →', size=14, color=SLATE, bold=True, align='c')
rx = 12.7; rw = X1 - rx
card(s, rx, 3.15, rw, 6.3)
icon_circle(s, rx + rw / 2, 3.15, 1.1, BLUE, 'gauge')
text(s, rx + 0.3, 3.95, rw - 0.6, 0.6, 'Plus one run-time factor', size=20, color=INK, bold=True, align='c')
text(s, rx + 0.4, 4.7, rw - 0.8, 4.6, [
    {'t': '**Confidence.** When the agent is unsure (an unclear reply, a borderline risk score, a rule that almost applies), it steps down one level automatically and asks.', 'size': 15},
    {'t': 'The level is a ceiling, not a target.', 'size': 15, 'color': BLUE, 'space_before': 10, 'bold': True},
    {'t': 'Decide the level per action, write it down, and review it in the weekly log review.', 'size': 15, 'space_before': 10}],
    color=INK)
notes(s, """
A simple grid that non-technical teams can apply in minutes. Two questions per action: how bad is a mistake, and can
it be undone? Low impact and easy to undo: let it run. High impact and hard to undo: a person decides, always. The two
middle cases get approval or supervision. Then the run-time twist: confidence. Even an autonomous action steps down
when the agent is unsure. Try it live: ask the room where 'offer a freed slot to the waiting list' belongs and why.
""")

# 25 ---- Trust ladder -----------------------------------------------------------------------
s = content_slide('Autonomy is earned, not granted', 'Run the same agent in stricter modes first, and loosen it as the evidence comes in')
stairs = [('Shadow', 'Agent proposes; nobody acts on it. Measure how often it was right.', 'Accuracy on real cases for 2–4 weeks', LIGHT, NAVY),
          ('Approve', 'Every action goes through a person.', 'Approval rate above target, no rule breaches', SKY, WHITE),
          ('Supervise', 'Agent acts; people handle exceptions.', 'Exception rate falling, outcomes stable', BLUE, WHITE),
          ('Autonomous', 'Agent acts within hard boundaries; people audit.', 'Ongoing audits, unchanged boundaries', NAVY, WHITE)]
sw2 = (XW - 3 * 0.25) / 4; base = 9.05
for i, (t, d, crit, col, tc) in enumerate(stairs):
    x = X0 + i * (sw2 + 0.25); top = 6.9 - i * 1.0; h = base - top
    b = rect(s, x, top, sw2, h, fill=col, radius=0.2, shadow=True)
    text(s, x + 0.3, top + 0.2, sw2 - 0.6, 0.55, t, size=22, color=tc, bold=True)
    text(s, x + 0.3, top + 0.8, sw2 - 0.6, h - 0.9, d, size=14, color=tc)
    # criteria to move up (above the step)
    text(s, x + 0.1, top - 1.35, sw2 - 0.2, 1.25, [{'t': 'Move up when' if i < 3 else 'Stay only while', 'bold': True, 'color': BLUE, 'size': 13},
                                                  {'t': crit, 'size': 13, 'color': INK}], anchor='b')
    if i < 3:
        chevron(s, x + sw2 - 0.12, top - 0.95, 0.5, 0.42, fill=BLUE, rotation=315)
text(s, X0, 2.95, 8.6, 2.2, [
    {'t': 'The No-Show Agent goes live in **Shadow** mode.', 'size': 18},
    {'t': 'Each step to the right is unlocked by evidence from the previous one, not by a calendar. '
          'Different clinics can run the same agent at different steps.', 'size': 15, 'color': SLATE, 'space_before': 6}],
    color=INK)
callout(s, X0, 9.4, XW, 0.85, 'You can always step down a level. That is a switch, not a project.', icon_name='sliders', size=16)
notes(s, """
The practical answer to 'how much autonomy should we give?': start strict and earn your way up with evidence. Shadow
mode costs nothing and produces the accuracy numbers. Approve mode builds trust with the coordinators who will later
supervise. Move up only when the numbers say so, and remember the switch goes both ways. Many organisations run the
same agent at different levels in different clinics; that is fine and expected.
""")

# 26 ---- Section: boundaries ----------------------------------------------------------------
s = section_slide('Boundaries and escalation', 'Designing the fence, and the door out', 'fence')
notes(s, "Section 5 of 5. About 7 minutes: four kinds of boundary, escalation triggers and a good hand-over, then the operating agreement on one page.")

# 27 ---- Boundaries -------------------------------------------------------------------------
s = content_slide('Boundaries: the agent\'s job description', 'Four kinds of limits, written down before the agent goes live')
bcols = [('layers', 'Scope', 'Which tasks, which patients, which clinics', ['Outpatient follow-ups in three clinics', 'Not oncology, not paediatrics', 'Logistics only, never clinical content']),
         ('key', 'Permissions', 'Which tools and data, read or write', ['Read appointments and contact details', 'Write bookings, within policy', 'No access to clinical notes or results']),
         ('gauge', 'Limits', 'How much, how often, when', ['Max two contact attempts per patient', 'Messages only between 08:00 and 20:00', 'Max 200 messages a day; pause above that']),
         ('stop', 'Hard rules', 'Must-never lines, no exceptions', ['Never cancel without consent', 'Never give medical advice', 'Always say it is an automated assistant'])]
bw = (XW - 3 * 0.4) / 4
for i, (ic, t, sub, bl) in enumerate(bcols):
    x = X0 + i * (bw + 0.4); y = 3.45; h = 6.7
    card(s, x, y, bw, h)
    icon_circle(s, x + bw / 2, y, 1.1, [NAVY, BLUE, SKY, RED][i], ic)
    text(s, x + 0.3, y + 0.8, bw - 0.6, 0.6, t, size=22, color=INK, bold=True, align='c')
    text(s, x + 0.3, y + 1.4, bw - 0.6, 0.8, sub, size=14, color=SLATE, align='c')
    text(s, x + 0.25, y + 2.35, bw - 0.5, 0.4, 'No-Show Agent', size=13, color=BLUE, bold=True)
    text(s, x + 0.35, y + 2.8, bw - 0.6, 3.7, [{'t': b, 'bullet': True} for b in bl], size=15, color=INK, space_after=8)
notes(s, """
Boundaries are the fence. Four kinds: scope (what it works on), permissions (what it can reach), limits (how much and
when), and hard rules (never, ever). All four are written in plain language by the process owners and enforced
technically by the runtime and the tool layer; the model does not get to argue with them. Show the No-Show Agent
examples and invite the room to add one more limit they would insist on.
""")

# 28 ---- Escalation -------------------------------------------------------------------------
s = content_slide('Escalation: when the agent must hand over', 'Escalation is a feature, not a failure')
lw3 = 8.55
card(s, X0, 3.0, lw3, 7.2)
heading(s, X0 + 0.45, 3.3, lw3 - 0.9, 'Triggers that stop the agent and call a person', size=19, color=BLUE)
trig = [('question', 'Low confidence', 'An unclear reply, a borderline risk score, a rule that almost applies'),
        ('stethoscope', 'Out of scope', '“Can I take a double dose until my appointment?” — any clinical question'),
        ('warning', 'High-risk action needed', 'The next step is irreversible or outside its permissions'),
        ('scale', 'Rule conflict', 'Two policies disagree, or the case is not covered by any'),
        ('repeat', 'Repeated failure', 'No reply after two attempts, or a tool is unavailable'),
        ('alarm', 'Patient asks for a human, or shows distress', 'Immediately, no questions asked')]
for i, (ic, a, b) in enumerate(trig):
    icon_row(s, X0 + 0.45, 3.95 + i * 1.02, lw3 - 0.9, ic, a, b, circle=NAVY if i % 2 == 0 else BLUE, d=0.56, label_size=15, desc_size=13, row_h=0.95)
rx = X0 + lw3 + 0.45; rw = X1 - rx
card(s, rx, 3.0, rw, 7.2)
heading(s, rx + 0.45, 3.3, rw - 0.9, 'What a good hand-over looks like', size=19, color=BLUE)
hrows = [('user', 'Goes to a named role', 'Clinic coordinator for logistics; nurse line for anything clinical'),
         ('files', 'Carries the full context', 'What happened, what the agent was about to do, and why it stopped'),
         ('clock', 'Has a time target', 'Clinical escalations answered within minutes, logistics within the day'),
         ('message', 'Tells the patient what happens next', '“A member of our team will contact you today.”')]
for i, (ic, a, b) in enumerate(hrows):
    icon_row(s, rx + 0.45, 3.95 + i * 1.55, rw - 0.9, ic, a, b, circle=BLUE, d=0.56, label_size=15, desc_size=13, row_h=1.4)
notes(s, """
The door in the fence. Six triggers; the last one is absolute. The right-hand card is what people forget: an escalation
is only useful if it lands with a named role, carries the context, has a time target, and the patient is told what
happens next. Ask the room: who would receive the escalations from the No-Show Agent in your clinic, and would they
have time? That question often decides whether a use case is realistic.
""")

# 29 ---- Operating agreement ----------------------------------------------------------------
s = content_slide('The No-Show Agent\'s operating agreement', 'One page that says what it may do alone, what needs a person, and who it calls')
rows_ = [('Action', 'Runs how', 'Boundary', 'Escalates to'),
         ('Send an appointment reminder', '4 · Autonomous', 'Approved templates only; 08:00–20:00; max two attempts', '—'),
         ('Rebook a routine follow-up', '4 · Autonomous, with patient consent', 'Within 14 days, same clinic, never cancel', 'Coordinator if no slot fits'),
         ('Offer a freed slot to the waiting list', '3 · Supervise', 'Same specialty, in priority order', 'Coordinator, daily summary'),
         ('Move a specialist appointment', '2 · Approve', 'Never without the clinician\'s OK', 'Clinic lead, same day'),
         ('Any clinical question or distress', 'Human only', 'The agent never answers; it acknowledges and hands over', 'Nurse line, immediately')]
tx, ty, tw_, th = X0, 3.0, XW, 6.55
tbl = s.shapes.add_table(len(rows_), 4, Inches(tx), Inches(ty), Inches(tw_), Inches(th)).table
colw = [4.6, 3.7, 5.35, 3.6]
for i, w_ in enumerate(colw): tbl.columns[i].width = Inches(w_)
tbl.rows[0].height = Inches(0.75)
for r in range(1, len(rows_)): tbl.rows[r].height = Inches((th - 0.75) / (len(rows_) - 1))
for r, row in enumerate(rows_):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (WHITE if r % 2 == 1 else ARC))
        cell.margin_left = Inches(0.2); cell.margin_right = Inches(0.2); cell.margin_top = Inches(0.08); cell.margin_bottom = Inches(0.08)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame; tf.word_wrap = True
        if r == 0:
            fill_paras(tf, val, size=16, color=WHITE, bold=True)
        else:
            bold = (c == 0)
            col = BLUE if c == 1 else INK
            if c == 1 and 'Human' in val: col = RED
            fill_paras(tf, val, size=14, color=col, bold=bold)
# strip table style effects
tblPr = tbl._tbl.tblPr
tblPr.set('firstRow', '1'); tblPr.set('bandRow', '0')
notes(s, """
Everything from sections 4 and 5 for one agent, on one page. Read it row by row. This is the artefact a clinic manager
signs before go-live, and the artefact an auditor asks for afterwards. Note how the levels differ per action: the same
agent is autonomous for reminders and 'human only' for anything clinical. Suggest the room keeps this format for
their own use cases: action, level, boundary, escalation route.
""")

# 30 ---- Exercise ---------------------------------------------------------------------------
s = content_slide('Your turn: where would you put the human?', '5 minutes in pairs. Three EHS scenarios, three quick decisions each.')
scen = [('flask', 'A · Lab results agent', 'Tells patients when routine results are normal, and books a follow-up when a result is abnormal.'),
        ('pill', 'B · Repeat prescription agent', 'Handles renewal requests from patients with stable chronic conditions and prepares them for sign-off.'),
        ('chat', 'C · Patient feedback agent', 'Reads complaints and compliments, routes them to the right department, and replies to the patient.')]
scw = (XW - 2 * 0.45) / 3
for i, (ic, t, d) in enumerate(scen):
    x = X0 + i * (scw + 0.45); y = 3.5; h = 4.15
    card(s, x, y, scw, h)
    icon_circle(s, x + scw / 2, y, 1.15, [NAVY, BLUE, SKY][i], ic)
    text(s, x + 0.3, y + 0.85, scw - 0.6, 0.6, t, size=20, color=INK, bold=True, align='c')
    text(s, x + 0.45, y + 1.65, scw - 0.9, 2.3, d, size=15, color=INK, align='c')
qs = [('1', 'Which autonomy level?', 'Assist, Approve, Supervise or Autonomous, per action'),
      ('2', 'One boundary you insist on', 'Scope, permission, limit or hard rule'),
      ('3', 'One escalation trigger', 'And who receives it')]
qw2 = (XW - 2 * 0.45) / 3
for i, (n, a, b) in enumerate(qs):
    x = X0 + i * (qw2 + 0.45); y = 8.1
    rect(s, x, y, qw2, 1.35, fill=PALE, radius=0.2)
    c = oval(s, x + 0.65, y + 0.675, 0.7, fill=BLUE); shape_text(c, n, size=20, color=WHITE, bold=True, margins=(0, 0, 0, 0))
    text(s, x + 1.2, y + 0.2, qw2 - 1.4, 0.45, a, size=16, color=NAVY, bold=True)
    text(s, x + 1.2, y + 0.65, qw2 - 1.4, 0.6, b, size=13, color=SLATE)
text(s, X0, 9.75, XW, 0.45, 'We will use your answers in the use-case assessment session that starts at 11:30.', size=14, color=SLATE)
notes(s, """
Five minutes, pairs, no writing needed beyond a sticky note. Each pair takes one scenario and answers three questions.
Debrief for two minutes: collect one answer per scenario. Expected patterns: A is 'Autonomous for normal results only
if policy allows, Approve or human for anything abnormal'; B is 'Assist or Approve, clinician signs'; C is 'Supervise
for routing, Approve for replies'. Disagreements are good; they show the boundaries need writing down. Carry the answers
into the use-case assessment session.
""")

# 31 ---- Takeaways --------------------------------------------------------------------------
s = prs.slides.add_slide(LAY['SAS - Blank - White'])
rect(s, 0, 0, 20, 11.25, fill=BLUE, radius=0)
rect(s, 14.2, 0, 5.8, 11.25, fill=NAVY, radius=0)
text(s, 1.37, 0.85, 12.0, 1.0, 'Takeaways', size=56, color=WHITE, bold=True, anchor='m')
tks = [('01', 'An agent is a system, not a model', 'Goal, planning, knowledge, tools, actions, and human oversight as a building block.'),
       ('02', 'The loop runs until the goal is met or a boundary is hit', 'Think, act, observe, repeat. Every exit is designed and logged.'),
       ('03', 'Autonomy is earned and bounded', 'Set the level per action, write the boundaries and escalation routes before go-live, and step up only with evidence.')]
for i, (n, a, b) in enumerate(tks):
    y = 2.9 + i * 2.35
    text(s, 1.37, y - 0.15, 2.2, 1.4, n, size=60, color=LIGHT, bold=True, anchor='t')
    text(s, 3.7, y, 9.9, 0.7, a, size=24, color=WHITE, bold=True)
    text(s, 3.7, y + 0.7, 9.9, 1.2, b, size=17, color=LIGHT)
icon(s, 'arrowRight', 'w', 15.2, 3.6, 0.9)
text(s, 15.2, 4.7, 4.2, 0.6, 'Up next  ·  11:30', size=16, color=LIGHT, bold=True)
text(s, 15.2, 5.3, 4.2, 2.2, 'Identifying & assessing Agentic AI use cases', size=26, color=WHITE, bold=True)
text(s, 15.2, 7.4, 4.2, 1.6, 'Bring the No-Show Agent canvas and your exercise answers.', size=15, color=LIGHT)
copyright_line(s, color=LIGHT)
notes(s, """
Three sentences to remember. If they remember only one: autonomy is earned and bounded. Then hand over to the
use-case assessment session, which starts from exactly the canvas and the operating agreement you have just seen.
""")

# 32 ---- Closing ----------------------------------------------------------------------------
s = prs.slides.add_slide(LAY['SAS - Closing'])
set_ph(s.shapes.title, 'Thank you!', size=72)
set_ph(placeholder(s, 11), 'Questions?', size=28)
notes(s, "Open questions for two or three minutes, then break into the use-case assessment session.")

prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides))
